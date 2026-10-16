"""Targeted FPR errata; preserve ZIP parts and run formatting outside replacements.

Backups and change log go to the supplied QA directory. Run once on originals.
"""
import argparse
import json
from pathlib import Path
from zipfile import ZipFile
import shutil
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
BASE = 'RQ1_Model_Training_xboost/RQ1_docx/'
FILES = [BASE+'Bao_Cao_Thuc_Nghiem_RQ1_XGBoost.docx',
         BASE+'Bao_Cao_Tong_Hop_Thuc_Nghiem_RQ1_RandomForest_XGBoost.docx',
         BASE+'Giai_Thich_Chuyen_Sau_Tham_So_Chi_So_XGBoost_vs_RF.docx',
         BASE+'Tai_lieu_ho_tro/Cam_Nang_Toan_Dien_XGBoost_KLTN.docx',
         'dataset/CICID-2018_processed/Bao_Cao_Tong_Hop_Dataset_3.5M_CICIDS2018.docx']


def replace(p, old, new):
    nodes = p.xpath('.//w:t', namespaces=NS)
    text = ''.join(n.text or '' for n in nodes)
    if old not in text:
        return False
    start = text.index(old)
    end = start + len(old)
    offset = 0
    inserted = False
    for node in nodes:
        value = node.text or ''
        right = offset + len(value)
        if right > start and offset < end:
            prefix = value[:max(0, start-offset)]
            suffix = value[max(0, end-offset):]
            node.text = prefix + (new if not inserted else '') + suffix
            node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            inserted = True
        offset = right
    return True


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--qa-dir', type=Path, required=True)
    parser.add_argument('--verify-only', action='store_true')
    args = parser.parse_args()
    log = []
    for rel in FILES:
        path = ROOT / rel
        backup = args.qa_dir / 'originals' / rel
        if args.verify_only:
            from docx import Document
            current, old = Document(path), Document(backup)
            assert [(len(t.rows), len(t.columns)) for t in current.tables] == [
                (len(t.rows), len(t.columns)) for t in old.tables]
            with ZipFile(path) as dst, ZipFile(backup) as src:
                assert dst.testzip() is None
                assert dst.namelist() == src.namelist()
                assert all(dst.read(n) == src.read(n) for n in src.namelist()
                           if n != 'word/document.xml')
            print('PASS structure and preserved parts:', rel)
            continue
        backup.parent.mkdir(parents=True, exist_ok=True)
        if backup.exists():
            if backup.read_bytes() != path.read_bytes():
                raise RuntimeError(f'Backup differs; refusing repeat edit: {backup}')
        else:
            shutil.copy2(path, backup)
        with ZipFile(path) as src:
            entries = [(i, src.read(i.filename)) for i in src.infolist()]
        xml = etree.fromstring(dict((i.filename, b) for i,b in entries)['word/document.xml'])
        changes = []
        for p in xml.xpath('//w:p', namespaces=NS):
            before = ''.join(p.xpath('.//w:t/text()', namespaces=NS))
            pairs = [
                ('FPR Benign (%) [Sub-1.2]', 'FNR Attack (%)'),
                ('5.7481% (FP=12071, TN=197929)', '0.2324% (FP=1139, TN=488861)'),
                ('FPR Benign = 5.75%', 'FPR Benign = 0.2324%'),
                ('FPR Benign chỉ 5.7481%', 'FPR Benign là 0.2324%'),
                ('FPR Benign chỉ 5.75%', 'FPR Benign là 0.2324%'),
                ('Giảm thiểu báo động giả tốt hơn (FPR Benign là 0.2324%)', 'FPR Benign là 0.2324% trên tập test (chưa đối chiếu FPR RF)'),
                ('tỷ lệ cảnh báo giả FPR lên tới 6.11%', 'tỷ lệ bỏ sót tấn công FNR khoảng 6.11%'),
                ('tỷ lệ báo động giả FPR trên Benign lên tới 6.11%', 'tỷ lệ bỏ sót tấn công FNR khoảng 6.11%'),
                ('FPR trên Benign cao (6.11%)', 'FNR trên Attack khoảng 6.11%'),
                ('tỷ lệ FPR trên Benign giảm xuống 5.7481%', 'FNR trên Attack là 5.7481%; FPR trên Benign là 0.2324%'),
                ('Tỷ lệ cảnh báo sai trên Benign giảm xuống mức tối ưu 5.7481% (Sub-RQ 1.2).', 'FPR trên Benign là 0.2324% (Sub-RQ 1.2); FNR trên Attack là 5.7481%.'),
                ('Tỷ lệ Báo động giả = 5.7481%', 'Tỷ lệ Báo động giả = 0.2324%'),
                ('FPR trên Benign (5.7481% - Sub-RQ 1.2)', 'FPR trên Benign (0.2324% - Sub-RQ 1.2)'),
                ('FPR = 5.7481%', 'FPR = 0.2324%'),
                ('Hạn chế báo động giả (FPR = 5.75%)', 'FPR trên tập test là 0.2324%'),
                ('-2.13% (Ít báo động giả hơn)', 'Chưa đối chiếu FPR RF'),
                ('XGBoost thấp hơn 2.13% (Giảm hơn 10,000 cảnh báo giả trên Benign)', 'Chưa đối chiếu FPR RF'),
                ('3. Về Tỷ lệ Báo động sai FPR trên Benign (5.75% so với 7.88%): XGBoost nhận diện chính xác 488,861 / 490,000 mẫu Benign (Recall đạt 99.77%), giảm đáng kể tỷ lệ cảnh báo giả so với Random Forest (Recall Benign của RF là 92.12%, tức có tới ~38,600 luồng sạch bị nghi ngờ là tấn công).',
                 '3. Về FPR trên Benign: XGBoost báo nhầm 1,139 / 490,000 flow Benign, tương ứng 0.2324%. Chưa đối chiếu ma trận RF nên chưa kết luận mức chênh lệch báo nhầm giữa hai mô hình.'),
                ('(thấp hơn đáng kể so với mức ≈ 7.88% của Random Forest), đảm bảo hệ thống NIDS hoạt động ổn định và tin cậy trong môi trường thực tế.', '(1,139 / 490,000 flow Benign). Đây là kết quả offline; chưa đủ để khẳng định độ tin cậy ngoài thực tế hoặc chênh lệch FPR với RF.'),
                ('đảm bảo phân biệt vững chắc giữa hành vi người dùng sạch và các cuộc tấn công mạng.', 'trên tập test. FNR là 5.7481%; cần đánh giá riêng trên traffic thực tế.'),
                ("Tỷ lệ FPR chỉ 5.75% đảm bảo hệ thống vừa nhạy bén vừa giữ được tính ổn định cho mạng nội bộ.", 'FPR đo được là 0.2324% trên tập test; không suy ra mức báo nhầm của mạng nội bộ thực tế.'),
                ("Giữ FPR ở mức thấp (5.75%) giúp ngăn ngừa triệt để hội chứng 'Tê liệt cảnh báo' (Alert Fatigue) trong trung tâm SOC.", 'FPR offline là 0.2324%; FNR là 5.7481%. Cần đánh giá tải cảnh báo trên traffic thực tế trước khi kết luận về SOC.'),
                ('giảm hơn 10,000 cảnh báo giả trên Benign', 'FPR XGBoost là 0.2324%; chưa đối chiếu FPR của RF'),
            ]
            # Bare percentages occur in comparison cells, not tuning values (which omit %).
            if before == '5.7481%':
                pairs.append(('5.7481%', '0.2324%'))
            if before in ('≈ 13.88%', '≈ 7.88%'):
                pairs.append((before, 'Chưa đối chiếu'))
            for old, new in pairs:
                replace(p, old, new)
            if 'Trong 490,000 gói tin sạch' in before:
                replace(p, 'Trong 490,000 gói tin sạch, chỉ có 1,139 gói bị nghi ngờ nhầm', 'Trong 490,000 flow Benign, có 1,139 flow bị báo nhầm')
            if 'triệt tiêu báo động giả' in before:
                replace(p, before, 'Tỷ lệ 70% Benign và 30% Attack là cấu hình lấy mẫu của nghiên cứu, không bảo đảm FPR thấp ngoài thực tế. Trên tập test đã lưu, XGBoost có FPR = 1,139 / 490,000 = 0.2324% và FNR = 12,071 / 210,000 = 5.7481%. Chưa có cơ sở khẳng định FPR < 0.1% từ tỷ lệ lấy mẫu này.')
            after = ''.join(p.xpath('.//w:t/text()', namespaces=NS))
            if before != after:
                changes.append(dict(before=before, after=after))
        if not changes:
            raise RuntimeError(f'No matches: {rel}')
        # Inline erratum near the opening, with inherited paragraph formatting.
        from copy import deepcopy
        anchor = xml.find('w:body/w:p', NS)
        note = etree.Element('{'+NS['w']+'}p')
        props = anchor.find('w:pPr', NS)
        if props is not None:
            props = deepcopy(props)
            for x in list(props):
                if etree.QName(x).localname in ('pStyle', 'pageBreakBefore', 'keepNext'):
                    props.remove(x)
            note.append(props)
        run = etree.SubElement(note, '{'+NS['w']+'}r')
        text = etree.SubElement(run, '{'+NS['w']+'}t')
        text.text = ('Đính chính ngày 16/10/2026: FPR đo Benign bị báo thành Attack (0.2324%); '
                     'FNR đo Attack bị bỏ sót thành Benign (5.7481%). Số liệu lấy từ ma trận test 700,000 flow. '
                     'Cột tuning cũ được đổi tên thành FNR; chưa có FPR từng cấu hình. '
                     'Không suy ra FPR live từ số liệu offline.')
        anchor.addnext(note)
        updated = etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
        with ZipFile(path, 'w') as dst:
            for info, data in entries:
                dst.writestr(info, updated if info.filename == 'word/document.xml' else data)
        with ZipFile(path) as dst, ZipFile(backup) as src:
            assert dst.testzip() is None
            assert dst.namelist() == src.namelist()
            assert all(dst.read(n) == src.read(n) for n in src.namelist() if n != 'word/document.xml')
        log.append(dict(file=rel, changes=changes))
        print(rel, len(changes), 'paragraphs corrected; other ZIP parts unchanged')
    if not args.verify_only:
        (args.qa_dir / 'changes.json').write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()
