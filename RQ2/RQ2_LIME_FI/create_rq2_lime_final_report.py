"""BUOC 6 - Tong hop bao cao Word RQ2_LIME_Ket_Qua_Final.docx (Sub-RQ 2.1-2.4).

Doc ket qua tu 5 buoc truoc (agreement, domain, stability/robustness, faithfulness) va
xuat mot bao cao Word hoan chinh, dung style giong RQ2_SHAP/create_shap_progress_report.py.

Chay tu goc du an (sau khi da chay du Buoc 1-5):
    python RQ2_LIME/create_rq2_lime_final_report.py
"""
from pathlib import Path
import argparse
import json

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
TEXT = "1F2933"
BASE = Path(__file__).resolve().parent


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, bold=False, size=None, color=TEXT):
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead:
        style_run(p.add_run(bold_lead), bold=True)
    style_run(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    style_run(r, bold=True, size=14 if level == 1 else 12, color="000000")
    return p


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(label), bold=True, size=font_size, color="FFFFFF")
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = widths[i]
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_run(p.add_run(str(value)), size=font_size)
            if row_idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, image_path, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(9)
    r = c.add_run(caption)
    style_run(r, size=9, color="4B5563")


def set_page_layout(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)


def latest_subdir(root):
    candidates = sorted([p for p in root.iterdir() if p.is_dir()])
    if not candidates:
        raise FileNotFoundError(f"Khong tim thay thu muc con trong {root}")
    return candidates[-1]


def fmt(x, nd=3):
    try:
        if pd.isna(x):
            return "NA"
        return f"{x:.{nd}f}"
    except (TypeError, ValueError):
        return str(x)


def make_chart_robustness(summary_overall_path, out_path):
    df = pd.read_csv(summary_overall_path)
    fig, ax = plt.subplots(figsize=(6.4, 3.4), dpi=150)
    ax.plot(df["noise_percent"], df["mean_jaccard5"], marker="o", color="#1F4E79", label="Mean Jaccard@5 truoc/sau nhieu")
    ax2 = ax.twinx()
    ax2.plot(df["noise_percent"], 100 * df["prediction_changed_rate"], marker="s", color="#C0392B",
              linestyle="--", label="Ty le doi predicted_label (%)")
    ax.set_xlabel("Muc nhieu thoi gian (%)")
    ax.set_ylabel("Mean Jaccard@5 (LIME truoc vs sau nhieu)", color="#1F4E79")
    ax2.set_ylabel("Ty le doi nhan (%)", color="#C0392B")
    ax.set_ylim(0, 1)
    lines1, labels1 = ax.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax.legend(lines1 + lines2, labels1 + labels2, loc="center left", fontsize=8)
    ax.set_title("Robustness LIME truoc nhieu thoi gian (tap con 289 flow)", fontsize=11)
    fig.tight_layout()
    fig.savefig(out_path)
    plt.close(fig)


def make_chart_winrate(summary_overall_path, out_path):
    df = pd.read_csv(summary_overall_path)
    ks = sorted(df["k"].unique())
    shap_wr = [df[(df.method == "shap") & (df.k == k)]["win_rate_percent"].iloc[0] for k in ks]
    lime_wr = [df[(df.method == "lime") & (df.k == k)]["win_rate_percent"].iloc[0] for k in ks]
    x = range(len(ks))
    width = 0.35
    fig, ax = plt.subplots(figsize=(6.0, 3.4), dpi=150)
    ax.bar([i - width / 2 for i in x], shap_wr, width, label="SHAP", color="#1F4E79")
    ax.bar([i + width / 2 for i in x], lime_wr, width, label="LIME", color="#C0392B")
    ax.set_xticks(list(x))
    ax.set_xticklabels([f"k={k}" for k in ks])
    ax.set_ylabel("Win rate (%)")
    ax.set_ylim(0, 105)
    for i, v in enumerate(shap_wr):
        ax.text(i - width / 2, v + 1, f"{v:.1f}", ha="center", fontsize=8)
    for i, v in enumerate(lime_wr):
        ax.text(i + width / 2, v + 1, f"{v:.1f}", ha="center", fontsize=8)
    ax.set_title("Faithfulness win-rate SHAP vs LIME - cohort 1.339 flow", fontsize=11)
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(out_path)
    plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, default=BASE.parent)
    parser.add_argument("--shap-cohort", type=Path, default=None)
    parser.add_argument("--lime-dir", type=Path, default=None)
    parser.add_argument("--agreement-dir", type=Path, default=None)
    parser.add_argument("--domain-dir", type=Path, default=None)
    parser.add_argument("--stability-dir", type=Path, default=None)
    parser.add_argument("--faithfulness-dir", type=Path, default=None)
    parser.add_argument("--output", type=Path, default=BASE.parent / "RQ2_LIME_Ket_Qua_Final.docx")
    args = parser.parse_args()

    root = args.project.resolve()
    shap_cohort = (args.shap_cohort or latest_subdir(root / "RQ2_SHAP" / "rq2_paired_cohort")).resolve()
    lime_dir = (args.lime_dir or latest_subdir(root / "RQ2_LIME" / "rq2_lime_paired")).resolve()
    agreement_dir = (args.agreement_dir or latest_subdir(root / "RQ2_LIME" / "rq2_agreement")).resolve()
    domain_dir = (args.domain_dir or latest_subdir(root / "RQ2_LIME" / "rq2_domain")).resolve()
    stability_dir = (args.stability_dir or latest_subdir(root / "RQ2_LIME" / "rq2_stability_robustness")).resolve()
    faithfulness_dir = (args.faithfulness_dir or latest_subdir(root / "RQ2_LIME" / "rq2_faithfulness")).resolve()

    manifest = pd.read_csv(shap_cohort / "manifest_main.csv")
    lime_info = json.loads((lime_dir / "run_info.json").read_text(encoding="utf-8"))
    lime_scores = pd.read_csv(lime_dir / "lime_sample_scores.csv")

    agreement_by_class = pd.read_csv(agreement_dir / "summary_by_class.csv")
    agreement_macro = pd.read_csv(agreement_dir / "summary_macro.csv").iloc[0]

    domain_by_class = pd.read_csv(domain_dir / "summary_by_class.csv")
    domain_macro = pd.read_csv(domain_dir / "summary_macro.csv").iloc[0]

    stability_by_class = pd.read_csv(stability_dir / "summary_stability_by_class.csv")
    robustness_overall = pd.read_csv(stability_dir / "summary_robustness_overall.csv")
    robustness_by_class = pd.read_csv(stability_dir / "summary_robustness_by_class.csv")
    stability_info = json.loads((stability_dir / "run_info.json").read_text(encoding="utf-8"))

    faith_overall = pd.read_csv(faithfulness_dir / "summary_overall.csv")
    faith_by_class = pd.read_csv(faithfulness_dir / "summary_by_class.csv")

    charts_dir = BASE / "report_charts"
    charts_dir.mkdir(exist_ok=True)
    robustness_chart = charts_dir / "robustness_lime.png"
    winrate_chart = charts_dir / "winrate_shap_vs_lime.png"
    make_chart_robustness(stability_dir / "summary_robustness_overall.csv", robustness_chart)
    make_chart_winrate(faithfulness_dir / "summary_overall.csv", winrate_chart)

    doc = Document()
    set_page_layout(doc)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    style_run(title.add_run("RQ2 - Ket qua LIME ghep cap voi SHAP (Sub-RQ 2.1-2.4)"), bold=True, size=19, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    style_run(subtitle.add_run(
        f"Cohort ghep cap {len(manifest)} flow, seed 42 | XGBoost 15 lop, 78 feature | Cap nhat theo du lieu da chay"),
        size=10, color="4B5563")

    # ---------------------------------------------------------------- 1
    add_heading(doc, "1. Tom tat dieu hanh")
    k5_row = faith_overall[faith_overall.k == 5]
    shap_wr5 = k5_row[k5_row.method == "shap"]["win_rate_percent"].iloc[0]
    lime_wr5 = k5_row[k5_row.method == "lime"]["win_rate_percent"].iloc[0]
    add_body(doc,
        f"Tren cung {len(manifest)} flow (XGBoost du doan dung, seed 42, toi da 100 flow/lop), Jaccard@5 trung binh "
        f"giua Top-5 SHAP va Top-5 LIME (macro-average khong trong so 15 lop) la {fmt(agreement_macro['jaccard_5_macro_mean_of_class_means'])}, "
        f"Spearman {fmt(agreement_macro['spearman_5_macro_mean_of_class_means'])} va signed agreement "
        f"{fmt(agreement_macro['signed_agreement_5_macro_mean_of_class_means'])} tren cac feature chung. Muc do tuong dong Top-5 o "
        f"muc thap-trung binh; hai phuong phap thuong khong chon dung cung mot bo 5 feature nhung khi chung chon "
        f"trung mot feature thi phan lon cung dau.",
        "Sub-RQ 2.1 (tuong dong): ")
    add_body(doc,
        f"Domain Precision@5 (doi chieu Top-5 voi bang expected-signals 15 lop) dat macro-average "
        f"{fmt(domain_macro['shap_precision_at_5_macro_mean_of_class_means'])} cho SHAP va "
        f"{fmt(domain_macro['lime_precision_at_5_macro_mean_of_class_means'])} cho LIME - ca hai deu THAP (duoi 20% "
        f"tren trung binh), cho thay Top-5 dua tren attribution so voi mo hinh thuong KHONG trung voi tap feature "
        f"ma con nguoi ky vong tu co che tan cong; day la ket qua can bao cao trung thuc, khong phai bang chung ca hai "
        f"phuong phap 'sai', ma la gioi han cua ca dinh nghia domain signal don gian va cua chinh Top-5.",
        "Sub-RQ 2.2 (domain validation): ")
    add_body(doc,
        f"LIME on dinh vua phai giua cac lan chay khac seed (median Jaccard@5 giua cac cap seed = "
        f"{fmt(pd.read_csv(stability_dir / 'per_flow_stability.csv')['stability_jaccard5_mean'].median())} tren tap con "
        f"{stability_info['n_subset']} flow) va giam dan do on dinh khi nhieu thoi gian tang (xem Hinh 1). Local R^2 cua "
        f"surrogate LIME THAP: median {fmt(lime_info['median_local_r2'])}, mean {fmt(lime_info['mean_local_r2'])} tren "
        f"toan bo {lime_info['n_flow']} flow - duoi nguong canh bao 0,3; chuan doan tren tap con 60 flow cho thay "
        f"tang num_samples len 10.000 hoac tat discretize_continuous KHONG cai thien R^2, nen cau hinh mac dinh "
        f"(num_samples=5.000, discretize_continuous=True) duoc giu lai va gioi han nay duoc ghi ro o Muc 6.",
        "Sub-RQ 2.3 (stability & robustness): ")
    add_body(doc,
        f"Voi phep thu donor-substitution giong het thiet ke SHAP full-test nhung chay tren CUNG {len(manifest)} flow "
        f"cho ca hai phuong phap: tai k=5, SHAP thang doi chung ngau nhien o {fmt(shap_wr5,2)}% flow du dieu kien, "
        f"LIME o {fmt(lime_wr5,2)}% - SHAP co win-rate cao hon LIME o ca ba muc k (xem Bang 5). Day la pham vi "
        f"{len(manifest)} flow, KHAC voi con so 99,31% da cong bo tren 700.000 flow cua SHAP goc; hai con so khong "
        f"duoc gop chung.",
        "Sub-RQ 2.4 (faithfulness, quan trong nhat): ")

    # ---------------------------------------------------------------- 2
    doc.add_page_break()
    add_heading(doc, "2. Sub-RQ 2.1 - Do tuong dong SHAP-LIME theo 15 lop")
    add_body(doc, "Jaccard@k = |Top-k SHAP giao Top-k LIME| / |hop|. Spearman va signed agreement chi tinh tren cac "
                  "feature CHUNG giua hai Top-k (NA neu duoi 2 feature chung). Bang duoi la k=5 (xem CSV day du cho "
                  "k=3 va k=10 trong summary_by_class.csv).")
    rows = []
    for _, r in agreement_by_class.iterrows():
        rows.append([r["true_label"], int(r["N"]), fmt(r["jaccard_5_mean"]), fmt(r["jaccard_5_median"]),
                     fmt(r["spearman_5_mean"]), fmt(r["signed_agreement_5_mean"]), int(r["spearman_5_n_valid"])])
    add_table(doc, ["Lop", "N", "Jaccard@5 mean", "Jaccard@5 median", "Spearman mean", "Signed agree. mean", "N Spearman huu ich"],
               rows, [Inches(1.55), Inches(0.4), Inches(0.85), Inches(0.9), Inches(0.85), Inches(0.95), Inches(0.9)])
    add_body(doc,
        f"Macro-average khong trong so tren 15 lop (k=5): Jaccard={fmt(agreement_macro['jaccard_5_macro_mean_of_class_means'])}, "
        f"Spearman={fmt(agreement_macro['spearman_5_macro_mean_of_class_means'])}, "
        f"Signed agreement={fmt(agreement_macro['signed_agreement_5_macro_mean_of_class_means'])}. "
        f"SQL Injection chi co N=9 flow - Jaccard@5 = 0 va Spearman khong xac dinh (khong du 2 feature chung o hau "
        f"het flow); day la mau nho, KHONG duoc suy rong ra ket luan chung cho lop nay.",
        "Luu y co ban: ")

    # ---------------------------------------------------------------- 3
    add_heading(doc, "3. Sub-RQ 2.2 - Domain Precision@5: SHAP vs LIME")
    add_body(doc, "Domain Precision@5 = (so feature trong Top-5 khop voi bang expected-signals cua LOP THAT) / 5, "
                  "tinh rieng cho SHAP va LIME tren CUNG cohort de so sanh cong bang.")
    rows = []
    for _, r in domain_by_class.iterrows():
        rows.append([r["true_label"], int(r["N"]), fmt(r["shap_precision_at_5_mean"]), fmt(r["lime_precision_at_5_mean"])])
    add_table(doc, ["Lop", "N", "SHAP Precision@5 mean", "LIME Precision@5 mean"], rows,
               [Inches(1.8), Inches(0.5), Inches(1.6), Inches(1.6)])
    add_body(doc,
        f"Macro-average: SHAP={fmt(domain_macro['shap_precision_at_5_macro_mean_of_class_means'])}, "
        f"LIME={fmt(domain_macro['lime_precision_at_5_macro_mean_of_class_means'])}. Ca hai deu thap; SHAP nhinh hon "
        f"LIME o hau het lop. Day la doi chieu djnh tinh co cau truc voi bang expected-signals suy tu co che tan "
        f"cong (Bang 8, ke hoach chi tiet), KHONG phai diem so chuyen gia SOC doc lap, nen khong duoc cong bo nhu "
        f"'ty le dung' cua XAI.",
        "Cach doc: ")

    # ---------------------------------------------------------------- 4
    doc.add_page_break()
    add_heading(doc, "4. Sub-RQ 2.3 - Stability va Robustness cua LIME")
    add_body(doc,
        f"Stability: chay lai LIME voi {len(stability_info['stability_seeds'])} seed khac nhau tren tap con "
        f"{stability_info['n_subset']} flow ({stability_info['n_per_class']} flow/lop). Voi moi flow, tinh Jaccard@5 "
        f"trung binh giua tat ca cap seed; bang duoi la median theo lop.")
    rows = []
    for _, r in stability_by_class.iterrows():
        rows.append([r["true_label"], int(r["N"]), fmt(r["median"]), fmt(r["mean"]), fmt(r["std"])])
    add_table(doc, ["Lop", "N", "Median Jaccard@5", "Mean", "Std"], rows,
               [Inches(1.8), Inches(0.5), Inches(1.2), Inches(1.2), Inches(1.2)])

    add_body(doc,
        f"Robustness: ap dung dung co che nhieu thoi gian cua SHAP (s = 1 + a*U, a in {{1%, 5%, 10%}}, 3 seed) tren "
        f"CUNG tap con {stability_info['n_subset']} flow - KHONG phai toan bo cohort hay 700k flow, do chi phi tinh "
        f"toan LIME. Lop giai thich giu co dinh la lop du doan TRUOC nhieu.",
        "")
    rows = []
    for _, r in robustness_overall.iterrows():
        rows.append([f"{r['noise_percent']:.0f}%", int(r["N"]), fmt(r["mean_jaccard5"]), fmt(r["median_jaccard5"]),
                     fmt(100 * r["prediction_changed_rate"], 2) + "%"])
    add_table(doc, ["Muc nhieu", "N (flow x seed)", "Mean Jaccard@5", "Median Jaccard@5", "Ty le doi predicted_label"],
               rows, [Inches(1.1), Inches(1.3), Inches(1.4), Inches(1.4), Inches(1.4)])
    add_figure(doc, robustness_chart,
        "Hinh 1. Do on dinh Top-5 LIME (Jaccard@5) va ty le doi nhan du doan khi tang muc nhieu thoi gian, tren tap con.")

    # ---------------------------------------------------------------- 5
    doc.add_page_break()
    add_heading(doc, "5. Sub-RQ 2.4 - Faithfulness ghep cap SHAP vs LIME (quan trong nhat)")
    add_body(doc,
        "Ap dung DUNG protocol donor-substitution cua shap_faithfulness_full_test.py (10 donor tu train, seed 173; "
        "20 lan doi chung ngau nhien, seed 2026; k = 3, 5, 10; nguong thang 1e-6) cho CA SHAP VA LIME tren CUNG "
        f"{len(manifest)} flow cua cohort.")
    rows = []
    for _, r in faith_overall.iterrows():
        rows.append([r["method"].upper(), int(r["k"]), int(r["N"]), int(r["n_eligible"]), fmt(r["coverage_percent"], 2) + "%",
                     fmt(r["win_rate_percent"], 2) + "%", fmt(r["mean_margin_drop_method"]), fmt(r["mean_margin_drop_random"])])
    add_table(doc, ["Phuong phap", "k", "N", "N du dieu kien", "Coverage", "Win rate", "Margin drop (method)", "Margin drop (random)"],
               rows, [Inches(0.85), Inches(0.4), Inches(0.5), Inches(0.9), Inches(0.85), Inches(0.85), Inches(1.1), Inches(1.1)])
    add_figure(doc, winrate_chart, "Hinh 2. Win-rate SHAP vs LIME (donor-substitution) tren cohort 1.339 flow, k=3/5/10.")
    add_body(doc,
        f"SHAP thang doi chung ngau nhien o ty le cao hon LIME tai ca ba k (k=3: SHAP "
        f"{fmt(faith_overall[(faith_overall.method=='shap')&(faith_overall.k==3)]['win_rate_percent'].iloc[0],2)}% "
        f"so LIME {fmt(faith_overall[(faith_overall.method=='lime')&(faith_overall.k==3)]['win_rate_percent'].iloc[0],2)}%; "
        f"k=10: SHAP {fmt(faith_overall[(faith_overall.method=='shap')&(faith_overall.k==10)]['win_rate_percent'].iloc[0],2)}% "
        f"so LIME {fmt(faith_overall[(faith_overall.method=='lime')&(faith_overall.k==10)]['win_rate_percent'].iloc[0],2)}%). "
        f"Luu y k=10 cua SHAP chi co coverage "
        f"{fmt(faith_overall[(faith_overall.method=='shap')&(faith_overall.k==10)]['coverage_percent'].iloc[0],2)}% vi mot so flow "
        f"khong co du 10 feature voi SHAP duong.",
        "Ket luan Sub-RQ 2.4: ")
    add_body(doc,
        "KHONG duoc gop ket qua nay voi 99,31% da cong bo truoc do cho SHAP tren 700.000 flow test set day du. Hai "
        "con so do TREN HAI PHAM VI KHAC NHAU (700k flow toan bo vs 1.339 flow cohort ghep cap dung de so sanh voi "
        "LIME) va co the KHONG bang nhau ve mat so hoc dan du cung mot phuong phap. 99,31% chi duoc trich dan lam "
        "bang chung global BO SUNG cho do tin cay cua SHAP, khong thay the bang so sanh truc tiep o day.",
        "Canh bao pham vi (KHONG duoc bo qua): ")

    # ---------------------------------------------------------------- 6
    doc.add_page_break()
    add_heading(doc, "6. Gioi han")
    add_table(doc, ["Gioi han", "Chi tiet"], [
        ["Co mau nho", "SQL Injection N=9 va Brute Force-XSS N=44 (trong tong 1.339 flow) qua nho de ket luan manh "
                        "o cap lop; cac chi so cua hai lop nay chi mang tinh minh hoa, khong dai dien thong ke."],
        ["LIME phu thuoc cau hinh", f"kernel_width=0,75*sqrt(78)~{6.6238:.4f} va num_samples=5.000 la lua chon co dinh "
                                     f"truoc. Local R^2 thap (median {fmt(lime_info['median_local_r2'])}) cho thay surrogate "
                                     "tuyen tinh cuc bo kho khop voi ham quyet dinh phi tuyen cua XGBoost tren khong "
                                     "gian 78 chieu; day la dac diem cua chinh phuong phap LIME tren du lieu nay, "
                                     "khong phai loi cau hinh."],
        ["Domain signals", "Bang expected-signals (Bang 8, ke hoach chi tiet) la kien thuc mien suy luan tu co che "
                            "tan cong, khong phai diem so tu chuyen gia SOC doc lap hay doi chieu PCAP thuc te."],
        ["Robustness LIME pham vi hep", "Chi chay tren tap con toi da 20 flow/lop (289 flow), khong phai toan bo "
                                          "cohort hay 700k flow nhu SHAP, do chi phi tinh toan LIME lon hon nhieu."],
        ["Khong tron cac loai ty le", "Jaccard, Domain Precision@5, win-rate faithfulness va Local R^2 do CAC KHIA "
                                       "CANH KHAC NHAU cua loi giai thich; khong duoc gop thanh mot con so duy nhat."],
    ], [Inches(1.7), Inches(4.85)], font_size=9.5)

    # ---------------------------------------------------------------- 7
    add_heading(doc, "7. Ket luan RQ2")
    add_body(doc,
        f"Tren cohort ghep cap {len(manifest)} flow (XGBoost du doan dung, seed 42, toi da 100 flow/lop, phan tang "
        f"15 lop), macro-average (khong trong so, 15 lop) cho Jaccard@5 SHAP-LIME la "
        f"{fmt(agreement_macro['jaccard_5_macro_mean_of_class_means'])} - muc tuong dong Top-5 THAP-TRUNG BINH. "
        f"Domain Precision@5 macro-average la {fmt(domain_macro['shap_precision_at_5_macro_mean_of_class_means'])} cho SHAP va "
        f"{fmt(domain_macro['lime_precision_at_5_macro_mean_of_class_means'])} cho LIME, ca hai deu thap khi doi chieu voi "
        f"bang expected-signals suy tu co che tan cong. Faithfulness (donor-substitution, k=5, CUNG {len(manifest)} "
        f"flow) cho win-rate {fmt(shap_wr5,2)}% voi SHAP va {fmt(lime_wr5,2)}% voi LIME - SHAP bam sat quyet dinh "
        f"cua XGBoost hon LIME tren cohort nay xet theo phep thu da chon. LIME co Local R^2 trung binh thap "
        f"({fmt(lime_info['mean_local_r2'])}) va do on dinh Top-5 giua cac seed o muc trung binh, giam khi nhieu "
        f"thoi gian tang.",
        "")
    add_body(doc,
        "Bon chi so - Jaccard (2.1), Domain Precision@5 (2.2), win-rate faithfulness (2.4) va Local R^2/on dinh "
        "(2.3) - do BON khia canh khac nhau cua chat luong giai thich va KHONG duoc gop thanh mot con so 'RQ2 dung "
        "X%' duy nhat. Ket luan co dieu kien: tren du lieu CSE-CIC-IDS2018 va mo hinh XGBoost da huan luyen, SHAP "
        "va LIME thuong khong chon trung Top-5 feature, ca hai deu co do phu hop domain thap khi doi chieu voi kien "
        "thuc tan cong o muc don gian hoa hien tai, va SHAP the hien faithfulness cao hon LIME trong pham vi phep "
        "thu donor-substitution da ap dung tren chinh cohort nay.",
        "Kết luận cuối cùng: ")

    doc.core_properties.title = "RQ2 - Ket qua LIME ghep cap voi SHAP"
    doc.core_properties.subject = "Sub-RQ 2.1-2.4: agreement, domain validation, stability/robustness, faithfulness"
    doc.core_properties.author = "Nhom nghien cuu"
    doc.save(args.output)
    print(f"HOAN TAT. Bao cao: {args.output}")


if __name__ == "__main__":
    main()
