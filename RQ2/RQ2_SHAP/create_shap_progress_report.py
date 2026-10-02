"""Create the SHAP-only progress report for RQ2.

Run from the repository root:
  <bundled-python> RQ2_SHAP/create_shap_progress_report.py
"""

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "Bao_cao_tien_do_SHAP_RQ2.docx"
SUMMARY_SOURCE = BASE / "Bao_cao_RQ_SHAP_XGBoost.docx"
TOP5_SOURCE = BASE / "Bao_cao_Top5_SHAP_tung_loai_tan_cong.docx"
SOURCE_ASSETS = BASE / "report_assets_from_two_source_docs"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
TEXT = "1F2933"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, bold=False, size=None, color=TEXT):
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead:
        style_run(p.add_run(bold_lead), bold=True)
    style_run(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    style_run(r, bold=True, size=14 if level == 1 else 12, color="000000")
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(label), bold=True, size=10, color="FFFFFF")
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = widths[i]
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_run(p.add_run(str(value)), size=10)
            if row_idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, image_path, caption, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(9)
    r = c.add_run(caption)
    style_run(r, size=9, color="4B5563")


def set_page_layout(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)


def create_report():
    global doc
    doc = Document()
    set_page_layout(doc)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("Báo cáo tiến độ triển khai SHAP cho RQ2")
    style_run(title_run, bold=True, size=21, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    style_run(subtitle.add_run("Dự án phát hiện xâm nhập CSE CIC IDS2018 | Cập nhật ngày 27 tháng 9 năm 2026"), size=10, color="4B5563")

    add_heading(doc, "1 Tóm tắt tiến độ")
    add_body(
        doc,
        "Giai đoạn SHAP Top k đã hoàn thành. Bộ kết quả hiện đã sẵn sàng làm dữ liệu đối chứng cố định cho LIME; vì vậy không cần chạy lại SHAP khi bắt đầu phần LIME.",
        "Kết luận: ",
    )
    add_table(
        doc,
        ["Hạng mục", "Trạng thái", "Kết quả hiện có"],
        [
            ["Chọn flow dùng chung", "Hoàn thành", "1.363 flow phân tầng, seed 42; tối đa 100 flow mỗi lớp"],
            ["Giải thích TreeSHAP", "Hoàn thành", "78 đặc trưng trên lớp do XGBoost dự đoán cho từng flow"],
            ["Xuất Top 1 Top 3 Top 5", "Hoàn thành", "1.363, 4.089 và 6.815 dòng xếp hạng"],
            ["Kiểm tra additivity", "Hoàn thành", "Sai số tuyệt đối trung bình 1,97e-6; lớn nhất 1,14e-5"],
            ["So sánh SHAP LIME", "Chưa bắt đầu", "Chờ LIME chạy trên đúng 1.363 flow và cùng lớp đầu ra"],
        ],
        [Inches(1.75), Inches(1.05), Inches(3.55)],
    )

    add_heading(doc, "2 Khung báo cáo tiến độ")
    add_body(doc, "Khung này dùng để báo cáo với giảng viên hoặc nhóm dự án. Phần SHAP chỉ báo cáo những việc đã chạy và tệp đã xuất; không tuyên bố RQ2 đã hoàn tất khi chưa có LIME.")
    add_table(
        doc,
        ["Mục", "Nội dung cần báo cáo"],
        [
            ["Mục tiêu", "Tạo bộ giải thích SHAP Top 1 Top 3 Top 5 để chuẩn bị so sánh công bằng với LIME."],
            ["Dữ liệu và mô hình", "Test set CIC IDS2018 đã xử lý; mô hình XGBoost cố định; 78 đặc trưng mạng."],
            ["Việc đã làm", "Chọn 1.363 flow dùng chung, giải thích từng flow, lưu ranking và kiểm tra additivity."],
            ["Kết quả minh họa", "Biểu đồ Top 5 theo lớp và waterfall của một flow được dự đoán đúng."],
            ["Ý nghĩa cho RQ2", "Đã khóa dữ liệu SHAP cho Sub RQ 2.1; có bằng chứng mô tả feature cho Sub RQ 2.2."],
            ["Việc tiếp theo", "Chạy LIME trên cùng flow, cùng feature name và cùng predicted label; sau đó tính Overlap và Jaccard."],
        ],
        [Inches(1.55), Inches(4.8)],
    )

    add_heading(doc, "3 Phạm vi SHAP đã thực hiện")
    add_body(doc, "Tập 1.363 flow được lấy theo phân tầng với seed 42. Các lớp có từ 100 flow trở lên được lấy 100 flow; lớp hiếm được lấy toàn bộ. Vì vậy có 46 flow XSS và 17 flow SQL Injection. Danh sách sample id này được khóa trong rq2_shared_samples.csv để LIME sử dụng lại.")
    add_table(
        doc,
        ["Chỉ số", "Giá trị", "Cách hiểu"],
        [
            ["Số flow SHAP", "1.363", "Bộ mẫu dùng chung SHAP và LIME, không phải toàn bộ test set."],
            ["Số feature", "78", "Tên feature và thứ tự được lưu cố định trong rq2_feature_schema.json."],
            ["Dự đoán đúng trong bộ mẫu", "1.152/1.363 (84,52%)", "Chỉ là bối cảnh của bộ mẫu SHAP; không phải accuracy RQ1 trên toàn test."],
            ["Cách giải thích", "TreeExplainer raw margin", "Giải thích contribution của feature vào lớp XGBoost đã dự đoán."],
        ],
        [Inches(1.55), Inches(1.35), Inches(3.45)],
    )
    add_body(doc, "Cách đọc điểm: giá trị SHAP dương đẩy raw score về lớp đang giải thích; giá trị âm kéo score ra xa lớp đó. SHAP không tự phân loại flow, mà giải thích quyết định đã có của XGBoost.")

    doc.add_page_break()
    add_heading(doc, "4 Kết quả SHAP đã xuất")
    add_body(doc, "Mỗi flow đã có ranking đủ 78 feature. Ba bảng Top k là phiên bản rút gọn của ranking này, phục vụ so sánh trực tiếp với LIME ở ba mức k = 1, 3 và 5.")
    add_table(
        doc,
        ["Tệp", "Vai trò trong RQ2"],
        [
            ["rq2_shared_samples.csv", "Khóa 1.363 sample id dùng chung; LIME bắt buộc dùng lại file này."],
            ["rq2_feature_schema.json", "Chuẩn hóa tên và thứ tự 78 feature để ghép SHAP với LIME."],
            ["rq2_shap_sample_predictions.csv", "Lưu true label, predicted label và trạng thái dự đoán đúng sai."],
            ["rq2_shap_top1.csv, top3.csv, top5.csv", "Danh sách Top k theo từng sample id; là đầu vào trực tiếp cho Overlap và Jaccard."],
            ["rq2_shap_top5_by_class.csv", "Tổng hợp mean absolute SHAP theo lớp trên các flow dự đoán đúng."],
            ["rq2_shap_additivity_check.csv", "Kiểm tra tổng SHAP gần bằng raw margin của XGBoost."],
        ],
        [Inches(2.35), Inches(4.0)],
    )

    add_figure(
        doc,
        FIGURES / "03_shap_top5_mean_by_class.png",
        "Hình 1. Top 5 mean absolute SHAP theo lớp trên các flow được XGBoost dự đoán đúng. Thanh dài hơn nghĩa là feature có mức ảnh hưởng trung bình lớn hơn.",
        width=6.35,
    )

    add_heading(doc, "5 Ví dụ giải thích một flow")
    add_body(doc, "Waterfall dưới đây là một flow LOIC HTTP có sample id 7748, được XGBoost dự đoán đúng với xác suất hiển thị gần 0,999. Các khối màu đỏ là feature đẩy raw score của lớp LOIC HTTP lên từ giá trị nền E[f(X)] đến score cuối f(x). Đây là ví dụ cục bộ cho một flow, không dùng để kết luận cho toàn bộ lớp.")
    add_figure(
        doc,
        FIGURES / "waterfall" / "waterfall_DDoS_attacks_LOIC_HTTP_sample_7748.png",
        "Hình 2. Waterfall SHAP cho flow LOIC HTTP sample id 7748. Fwd Pkts/s, Bwd Pkt Len Std và Fwd IAT Std là các đóng góp tăng điểm nổi bật trong flow này.",
        width=6.2,
    )

    doc.add_page_break()
    add_heading(doc, "6 Mức độ đáp ứng RQ2 sau giai đoạn SHAP")
    add_table(
        doc,
        ["Câu hỏi", "Trạng thái sau SHAP", "Điều còn thiếu"],
        [
            ["Sub RQ 2.1 Tương đồng SHAP LIME", "SHAP đã sẵn sàng làm chuẩn so sánh.", "LIME phải xuất Top 1 Top 3 Top 5 trên cùng 1.363 flow."],
            ["Sub RQ 2.2 Domain validation", "Đã có Top 5 theo lớp để đối chiếu với đặc điểm lưu lượng tấn công.", "Cần đối chiếu feature LIME và kết luận có điều kiện; không suy ra quan hệ nhân quả."],
            ["Sub RQ 2.3 Faithfulness", "Giai đoạn Top k này đã kiểm tra additivity số học.", "Nếu RQ2 yêu cầu so sánh hai phương pháp, cần chạy phép perturbation faithfulness tương tự cho LIME."],
        ],
        [Inches(1.85), Inches(2.2), Inches(2.3)],
    )
    add_body(doc, "Lưu ý quan trọng: additivity là kiểm tra tổng các contribution SHAP khớp gần với raw margin của XGBoost. Nó không phải phần trăm lời giải thích đúng, không thay thế faithfulness perturbation và không chứng minh cơ chế tấn công ngoài thực tế.")

    add_heading(doc, "7 Công việc tiếp theo cho LIME")
    add_table(
        doc,
        ["Bước", "Yêu cầu bắt buộc"],
        [
            ["1", "Nạp rq2_shared_samples.csv, lấy đúng 1.363 flow theo sample id."],
            ["2", "Dùng cùng mô hình XGBoost, cùng 78 feature và cùng thứ tự feature."],
            ["3", "Giải thích đúng predicted label đã lưu trong rq2_shap_sample_predictions.csv, không đổi sang true label."],
            ["4", "Xuất Lime Top 1 Top 3 Top 5 với feature_key giống file schema."],
            ["5", "Ghép từng sample id với SHAP và tính Overlap at k cùng Jaccard at k cho k bằng 1, 3, 5."],
            ["6", "Tổng hợp theo từng loại attack và chạy faithfulness cho LIME nếu muốn trả lời đầy đủ Sub RQ 2.3 cho cả hai phương pháp."],
        ],
        [Inches(0.6), Inches(5.75)],
    )

    add_heading(doc, "8 Kết luận tiến độ")
    add_body(doc, "Phần SHAP Top k đã hoàn thành và đã tạo đủ dữ liệu đầu vào để so sánh công bằng với LIME. Kết quả hiện tại chưa cho phép công bố phần trăm tương đồng SHAP LIME vì LIME chưa chạy. Sau khi có output LIME, nhóm chỉ cần ghép theo sample id và feature_key để hoàn thành Sub RQ 2.1, sau đó bổ sung faithfulness LIME để hoàn thiện Sub RQ 2.3 cho cả hai phương pháp.")

    doc.core_properties.title = "Báo cáo tiến độ triển khai SHAP cho RQ2"
    doc.core_properties.subject = "Tiến độ SHAP và kế hoạch tích hợp LIME cho RQ2"
    doc.core_properties.author = "Nhóm nghiên cứu"
    doc.save(OUTPUT)
    print(OUTPUT)


def extract_source_asset(source_doc, archive_member, output_name):
    """Copy an image embedded in one of the two approved source reports."""
    SOURCE_ASSETS.mkdir(parents=True, exist_ok=True)
    output_path = SOURCE_ASSETS / output_name
    with ZipFile(source_doc) as archive:
        output_path.write_bytes(archive.read(archive_member))
    return output_path


def create_source_only_report():
    """Build the progress report from the two user-approved DOCX reports only."""
    global doc
    if not SUMMARY_SOURCE.exists() or not TOP5_SOURCE.exists():
        raise FileNotFoundError("Không tìm thấy một trong hai báo cáo SHAP nguồn.")

    # image13 in the summary report is the full-test faithfulness comparison.
    faithfulness_chart = extract_source_asset(
        SUMMARY_SOURCE, "word/media/image13.png", "faithfulness_topk_from_summary_report.png"
    )
    # image29 in the Top-5 report is the positive Top-5 chart for LOIC-HTTP.
    loic_http_chart = extract_source_asset(
        TOP5_SOURCE, "word/media/image29.png", "loic_http_top5_from_top5_report.png"
    )

    doc = Document()
    set_page_layout(doc)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    style_run(title.add_run("Báo cáo tiến độ SHAP cho RQ2"), bold=True, size=21, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    style_run(
        subtitle.add_run(
            "Tổng hợp từ hai báo cáo SHAP đã có | Cập nhật ngày 27 tháng 9 năm 2026"
        ),
        size=10,
        color="4B5563",
    )

    add_heading(doc, "1 Tóm tắt tiến độ")
    add_body(
        doc,
        "Phần SHAP đã có kết quả giải thích theo lớp, Top 5 feature theo 14 loại tấn công, kiểm tra faithfulness có đối chứng và các hình minh họa. RQ2 chưa hoàn tất vì chưa có LIME trên cùng flow và cùng lớp đầu ra.",
        "Kết luận: ",
    )
    add_table(
        doc,
        ["Hạng mục", "Trạng thái", "Bằng chứng từ hai báo cáo"],
        [
            ["SHAP theo lớp", "Hoàn thành", "XGBoost 15 lớp, 78 feature; đã giải thích cục bộ và tổng hợp theo lớp."],
            ["Top 5 theo attack", "Hoàn thành", "14 lớp tấn công, 14 bảng feature và 28 biểu đồ; xử lý toàn bộ 700.000 flow test."],
            ["Đối chiếu domain", "Đã có bước đầu", "Nhận xét định tính về feature; chưa có chấm điểm chuyên gia hoặc PCAP log."],
            ["Faithfulness SHAP", "Hoàn thành", "Top 5 thắng đối chứng ngẫu nhiên ở 99,31% flow đủ điều kiện."],
            ["Tương đồng SHAP LIME", "Chưa có", "Chưa có LIME theo cùng flow và cùng đầu ra lớp."],
        ],
        [Inches(1.65), Inches(1.15), Inches(3.55)],
    )

    add_heading(doc, "2 Khung báo cáo tiến độ")
    add_body(doc, "Đây là khung ngắn có thể dùng khi báo cáo với giảng viên: mục tiêu, phạm vi, việc đã làm, kết quả minh họa, mức độ trả lời RQ2 và công việc còn lại.")
    add_table(
        doc,
        ["Mục", "Nội dung cần trình bày"],
        [
            ["Mục tiêu", "Dùng SHAP để biết XGBoost dựa vào feature nào khi nhận diện từng loại lưu lượng."],
            ["Phạm vi", "CSE CIC IDS2018, mô hình XGBoost cố định, 78 feature; báo cáo Top 5 bao phủ 14 lớp attack."],
            ["Việc đã làm", "Tính TreeSHAP, tổng hợp mean absolute và mean signed SHAP, xuất bảng và biểu đồ theo lớp."],
            ["Kết quả", "Feature nổi bật khác nhau giữa các lớp; có minh họa LOIC HTTP và phép faithfulness toàn test."],
            ["Ý nghĩa RQ2", "SHAP đã có bằng chứng cho domain validation định tính và faithfulness; chưa có similarity với LIME."],
            ["Việc tiếp theo", "Chạy LIME trên cùng flow, cùng output class và cùng k; sau đó tính Overlap và Jaccard."],
        ],
        [Inches(1.55), Inches(4.8)],
    )

    add_heading(doc, "3 Phạm vi và cách đọc kết quả SHAP")
    add_body(doc, "Hai báo cáo dùng TreeExplainer với model output raw, feature perturbation tree path dependent và không xấp xỉ. Báo cáo tổng hợp có một tập 1.363 flow phân tầng cho khảo sát và benchmark; báo cáo Top 5 tính lại SHAP trên đủ 700.000 flow test, trong đó phân tích riêng 14 lớp attack với tổng 210.000 flow.")
    add_body(doc, "Mean absolute SHAP cho biết mức ảnh hưởng trung bình của feature, không xét dấu. Mean signed SHAP cho biết feature trung bình đẩy điểm lớp lên hay kéo điểm lớp xuống. Đơn vị là raw margin, không phải phần trăm xác suất.")
    add_table(
        doc,
        ["Nhóm attack", "Feature Top 1", "Nhận xét đúng mức"],
        [
            ["Bot", "Dst Port 7,43631", "Model phụ thuộc mạnh vào cổng đích; đây có thể là dấu hiệu dịch vụ hoặc môi trường."],
            ["LOIC UDP", "Tot Fwd Pkts 9,21273", "Số gói chiều đi là feature chi phối trong dữ liệu test."],
            ["Slowloris", "Bwd IAT Max 7,66221", "Có liên hệ định tính với đặc tính thời gian và duy trì kết nối."],
            ["Infilteration", "Fwd Header Len 0,480182", "Feature flow không tái dựng đầy đủ chuỗi xâm nhập."],
            ["SQL Injection", "Init Fwd Win Byts 1,47507", "Không trực tiếp cho thấy payload SQL; lớp chỉ có 17 flow."],
        ],
        [Inches(1.15), Inches(1.65), Inches(3.55)],
    )

    doc.add_page_break()
    add_heading(doc, "4 Kết quả minh họa Top 5 SHAP")
    add_body(doc, "Hình dưới là biểu đồ gốc trong báo cáo Top 5. Với lớp DDoS LOIC HTTP, năm feature có tác động ròng dương cao nhất là Fwd Pkts/s, Fwd IAT Std, Init Fwd Win Byts, Bwd Pkt Len Std và Flow IAT Min. Biểu đồ này mô tả giá trị trung bình theo 51.354 flow mang nhãn thật LOIC HTTP, gồm cả dự đoán đúng và sai.")
    add_figure(
        doc,
        loic_http_chart,
        "Hình 1. Top 5 feature có tác động ròng dương của DDoS attacks LOIC HTTP. Nguồn: Báo cáo Top 5 SHAP theo từng loại tấn công.",
        width=6.25,
    )
    add_body(doc, "Ý nghĩa domain validation: tốc độ gói và khoảng thời gian giữa các gói là những tín hiệu lưu lượng có thể liên hệ với cường độ DDoS. Tuy nhiên, biểu đồ không tự xác nhận nguyên nhân tấn công, không cung cấp ngưỡng quyết định và không thay thế PCAP hoặc log dịch vụ.")

    add_heading(doc, "5 Mức độ trả lời RQ2")
    add_table(
        doc,
        ["Sub RQ", "Kết quả SHAP hiện có", "Cách báo cáo đúng"],
        [
            ["2.1 Similarity", "Chưa có tỷ lệ SHAP LIME.", "Không dùng stability giữa các lần SHAP thay cho similarity SHAP LIME."],
            ["2.2 Domain validation", "Có Top 5 và nhận xét theo từng attack.", "Kết luận định tính, không công bố phần trăm phù hợp thực tế."],
            ["2.3 Faithfulness", "Có phép thay feature với đối chứng ngẫu nhiên.", "Báo cáo rõ k, mẫu số, coverage và giới hạn can thiệp."],
        ],
        [Inches(1.25), Inches(2.45), Inches(2.65)],
    )
    add_body(doc, "Kết quả faithfulness Top 5: thay Top 5 feature SHAP làm raw margin của lớp dự đoán ban đầu giảm mạnh hơn thay ngẫu nhiên ở 695.140 trên 699.996 flow đủ điều kiện, tương đương 99,31%. Đây là win rate của phép thử, không phải accuracy của SHAP hay bằng chứng nhân quả.")
    add_figure(
        doc,
        faithfulness_chart,
        "Hình 2. Mức giảm raw margin khi thay các feature Top k SHAP và feature ngẫu nhiên. Nguồn: Báo cáo kết quả nghiên cứu SHAP trên XGBoost.",
        width=5.8,
    )

    doc.add_page_break()
    add_heading(doc, "6 Công việc tiếp theo cho LIME")
    add_table(
        doc,
        ["Bước", "Yêu cầu để so sánh công bằng"],
        [
            ["1", "Dùng cùng flow x và cùng output class c mà SHAP đã giải thích."],
            ["2", "Dùng cùng tên feature gốc và cùng quy tắc chọn Top k; không ép đủ k bằng hệ số bằng không."],
            ["3", "Xuất Top 1, Top 3 và Top 5 LIME cho từng flow."],
            ["4", "Tính Overlap at k = 100 x số feature chung chia k."],
            ["5", "Tính Jaccard at k = số feature chung chia số feature thuộc hợp của hai Top k."],
            ["6", "Nếu so sánh faithfulness cả hai phương pháp, chạy phép can thiệp LIME tương tự SHAP."],
        ],
        [Inches(0.6), Inches(5.75)],
    )
    add_body(doc, "Sau khi có LIME, mới có thể trả lời Sub RQ 2.1 bằng phần trăm overlap hoặc Jaccard. Khi đó RQ2 sẽ có ba phần rõ ràng: độ tương đồng SHAP LIME, đối chiếu domain cho cả hai, và faithfulness của từng phương pháp.")

    add_heading(doc, "7 Kết luận tiến độ")
    add_body(doc, "SHAP đã hoàn thành phần việc độc lập: xác định feature quan trọng theo lớp, tạo biểu đồ Top 5, đối chiếu domain ở mức định tính và đo faithfulness có đối chứng. Phần còn thiếu duy nhất để so sánh SHAP với LIME là kết quả LIME được tạo trên cùng mẫu và cùng lớp đầu ra. Báo cáo này chỉ sử dụng hai tài liệu nguồn: Bao_cao_RQ_SHAP_XGBoost.docx và Bao_cao_Top5_SHAP_tung_loai_tan_cong.docx.")

    doc.core_properties.title = "Báo cáo tiến độ SHAP cho RQ2"
    doc.core_properties.subject = "Tổng hợp từ hai báo cáo SHAP nguồn"
    doc.core_properties.author = "Nhóm nghiên cứu"
    doc.save(OUTPUT)
    print(OUTPUT)


def create_full_source_only_report():
    """Create the fuller RQ2 progress report from the two approved SHAP reports."""
    global doc
    if not SUMMARY_SOURCE.exists() or not TOP5_SOURCE.exists():
        raise FileNotFoundError("Không tìm thấy một trong hai báo cáo SHAP nguồn.")

    faithfulness_chart = extract_source_asset(
        SUMMARY_SOURCE, "word/media/image13.png", "faithfulness_topk_from_summary_report.png"
    )
    waterfall_chart = extract_source_asset(
        SUMMARY_SOURCE, "word/media/image7.png", "waterfall_web_to_benign_from_summary_report.png"
    )
    beeswarm_chart = extract_source_asset(
        SUMMARY_SOURCE, "word/media/image10.png", "beeswarm_slow_http_test_from_summary_report.png"
    )
    loic_http_chart = extract_source_asset(
        TOP5_SOURCE, "word/media/image29.png", "loic_http_top5_from_top5_report.png"
    )

    doc = Document()
    set_page_layout(doc)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    style_run(title.add_run("Báo cáo tiến độ SHAP cho RQ2"), bold=True, size=21, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    style_run(
        subtitle.add_run(
            "Tổng hợp từ hai báo cáo SHAP nguồn | Cập nhật ngày 27 tháng 9 năm 2026"
        ),
        size=10,
        color="4B5563",
    )

    add_heading(doc, "1 Mục đích của RQ2")
    add_body(
        doc,
        "RQ2 không đo lại accuracy của XGBoost. Mục tiêu là kiểm tra lời giải thích XAI có đáng tin và hữu ích đến đâu: SHAP và LIME có chỉ ra những feature giống nhau không, các feature đó có hợp lý với cơ chế tấn công mạng không, và việc can thiệp các feature quan trọng có làm thay đổi quyết định của model không.",
    )
    add_table(
        doc,
        ["Sub RQ", "Câu hỏi cần trả lời", "Ý nghĩa"],
        [
            ["2.1 Similarity", "Top k feature của SHAP và LIME trùng nhau bao nhiêu phần trăm trên từng flow hoặc từng attack?", "Đo mức đồng thuận giữa hai phương pháp XAI."],
            ["2.2 Domain validation", "Feature SHAP và LIME nêu ra có liên hệ hợp lý với đặc điểm lưu lượng của tấn công không?", "Kiểm tra lời giải thích có ý nghĩa kỹ thuật, không chỉ là con số."],
            ["2.3 Faithfulness", "Feature được XAI chọn có thực sự ảnh hưởng đến score của model không?", "Kiểm tra lời giải thích bám vào quyết định XGBoost."],
        ],
        [Inches(1.25), Inches(3.0), Inches(2.1)],
    )
    add_body(doc, "Vì vậy, không được gộp các khái niệm này thành một tỷ lệ duy nhất. Accuracy mô hình, similarity SHAP LIME, domain validity và faithfulness đo các khía cạnh khác nhau.")

    add_heading(doc, "2 Tóm tắt tiến độ hiện tại")
    add_table(
        doc,
        ["Hạng mục SHAP", "Trạng thái", "Kết quả đã có"],
        [
            ["Giải thích theo lớp", "Hoàn thành", "TreeSHAP cho XGBoost 15 lớp và 78 feature mạng."],
            ["Top 5 theo attack", "Hoàn thành", "14 lớp attack; 14 bảng feature, 28 biểu đồ và 14 ảnh kết quả chạy."],
            ["Phân tích đúng sai", "Hoàn thành", "So sánh đóng góp SHAP của nhóm dự đoán đúng và sai ở các lớp khó."],
            ["Domain validation", "Đã thực hiện bước đầu", "Đối chiếu định tính feature với hành vi tấn công; chưa có phần trăm do chưa chấm bởi chuyên gia."],
            ["Faithfulness", "Hoàn thành", "Thử thay Top k feature SHAP bằng giá trị train và so sánh với feature ngẫu nhiên."],
            ["Similarity SHAP LIME", "Chưa hoàn thành", "Chưa có LIME trên cùng flow và cùng output class."],
        ],
        [Inches(1.65), Inches(1.25), Inches(3.45)],
    )

    add_heading(doc, "3 Phạm vi và phương pháp SHAP đã làm")
    add_body(doc, "Báo cáo tổng hợp sử dụng mô hình XGBoost với n estimators 100, max depth 8, learning rate 0,08 và tree method hist. TreeExplainer được cấu hình model output raw, feature perturbation tree path dependent và không dùng xấp xỉ. Mỗi lời giải thích gồm giá trị nền và đóng góp của 78 feature đối với một output class.")
    add_body(doc, "Đã dùng 1.363 flow phân tầng cho khảo sát và benchmark. Riêng báo cáo Top 5 tính lại SHAP trên đủ 700.000 flow test. Trong đó, 14 lớp attack có 210.000 flow được phân tích riêng; 490.000 flow Benign đã được xử lý nhưng không có mục Top 5 riêng trong báo cáo đó.")
    add_body(doc, "Mean absolute SHAP là mức quan trọng trung bình của feature. Mean signed SHAP là chiều tác động trung bình: dương đẩy raw margin về lớp đang giải thích, âm kéo raw margin ra xa lớp đó. Raw margin là score trước chuyển sang xác suất; nó không phải phần trăm.")

    doc.add_page_break()
    add_heading(doc, "4 SHAP đáp ứng Sub RQ 2.2 như thế nào")
    add_body(doc, "SHAP đã tạo đủ bằng chứng để đối chiếu feature quan trọng với kiến thức mạng ở mức định tính. Bảng dưới tóm tắt feature đứng đầu nhóm tác động dương của đủ 14 attack trong báo cáo Top 5.")
    add_table(
        doc,
        ["Attack", "Feature Top 1 dương", "Mean absolute SHAP", "Ý nghĩa và giới hạn"],
        [
            ["Bot", "Dst Port", "7,43631", "Dấu hiệu dịch vụ hoặc môi trường, chưa xác nhận điều khiển bot."],
            ["Brute Force Web", "Fwd Pkt Len Mean", "2,55048", "Cần log yêu cầu và đăng nhập để xác nhận thử lặp."],
            ["Brute Force XSS", "Init Fwd Win Byts", "2,51924", "Không trực tiếp thể hiện mã script XSS."],
            ["HOIC", "Init Fwd Win Byts", "6,47441", "Dấu hiệu kết nối mạnh trong dữ liệu, chưa chứng minh cơ chế DDoS."],
            ["LOIC UDP", "Tot Fwd Pkts", "9,21273", "Số gói chiều đi là dấu hiệu lưu lượng cần kiểm tra thêm."],
            ["LOIC HTTP", "Fwd Pkts/s", "1,23100", "Tốc độ gói liên hệ với cường độ lưu lượng, không cho ngưỡng tấn công."],
            ["GoldenEye", "Fwd Seg Size Min", "8,18840", "Dấu hiệu segment, chưa đủ đặc hiệu để khẳng định GoldenEye."],
            ["Hulk", "Fwd Seg Size Min", "6,49332", "Dấu hiệu segment mạnh, không tự chứng minh cơ chế DoS."],
            ["SlowHTTPTest", "Fwd Seg Size Min", "5,09702", "Cần đọc cùng Flow Duration và Flow IAT Max."],
            ["Slowloris", "Bwd IAT Max", "7,66221", "Liên hệ định tính với thời gian duy trì kết nối."],
            ["FTP BruteForce", "Fwd Seg Size Min", "6,03244", "Không trực tiếp quan sát được số lần xác thực thất bại."],
            ["Infilteration", "Fwd Header Len", "0,480182", "Header không tái dựng đầy đủ chuỗi xâm nhập."],
            ["SQL Injection", "Init Fwd Win Byts", "1,47507", "Không có payload SQL; lớp chỉ có 17 flow."],
            ["SSH Bruteforce", "Dst Port", "6,20788", "Cần log xác thực để xác nhận hành vi brute force."],
        ],
        [Inches(1.2), Inches(1.45), Inches(1.0), Inches(2.1)],
    )
    add_body(doc, "Kết luận Sub RQ 2.2: có các liên hệ hợp lý, ví dụ LOIC UDP nổi bật bởi số gói chiều đi và Slowloris nổi bật bởi feature thời gian. Tuy nhiên, các feature flow không trực tiếp nhìn thấy payload XSS hoặc SQL Injection. Hai báo cáo chưa có bộ tiêu chí chấm điểm, chuyên gia độc lập hay PCAP log tương ứng, nên không được công bố mức phù hợp thực tế X phần trăm.")

    add_figure(
        doc,
        loic_http_chart,
        "Hình 1. Top 5 feature tác động ròng dương của DDoS attacks LOIC HTTP trên 51.354 flow. Nguồn: Báo cáo Top 5 SHAP theo từng loại tấn công.",
        width=6.25,
    )

    add_heading(doc, "5 Hình SHAP dùng để báo cáo")
    add_body(doc, "Biểu đồ Top 5 phù hợp để xem xu hướng theo một attack. Waterfall phù hợp để giải thích một cảnh báo cụ thể từ giá trị nền đến score cuối. Beeswarm phù hợp để xem phân bố đóng góp của nhiều flow; màu biểu thị giá trị feature, không biểu thị dự đoán đúng hay sai.")
    add_figure(
        doc,
        waterfall_chart,
        "Hình 2. Waterfall của một flow Brute Force Web có nhãn thật Web nhưng được giải thích ở output Benign. Nguồn: Báo cáo kết quả nghiên cứu SHAP trên XGBoost.",
        width=5.85,
    )
    add_figure(
        doc,
        beeswarm_chart,
        "Hình 3. Beeswarm cho lớp DoS attacks SlowHTTPTest trên tập 100 flow phân tầng. Nguồn: Báo cáo kết quả nghiên cứu SHAP trên XGBoost.",
        width=5.95,
    )

    doc.add_page_break()
    add_heading(doc, "6 SHAP đáp ứng Sub RQ 2.3 như thế nào")
    add_body(doc, "Faithfulness được kiểm tra bằng cách giữ cố định lớp XGBoost dự đoán ban đầu, chọn Top k feature có SHAP dương, rồi thay giá trị của chúng bằng giá trị từ 10 dòng train. Đối chứng chọn ngẫu nhiên k feature trong cùng tập feature SHAP dương, lặp 20 lần. Một flow được tính SHAP thắng khi raw margin giảm bởi Top k SHAP lớn hơn mức giảm bởi đối chứng ngẫu nhiên.")
    add_table(
        doc,
        ["k", "Flow đủ điều kiện", "Coverage", "Win rate SHAP"],
        [
            ["3", "700.000", "100,0000%", "98,17%"],
            ["5", "699.996", "99,9994%", "99,31%"],
            ["10", "649.956", "92,8509%", "99,83%"],
        ],
        [Inches(0.8), Inches(1.8), Inches(1.4), Inches(1.8)],
    )
    add_body(doc, "Với Top 5, raw margin giảm trung bình 3,4816 khi thay feature SHAP, so với 1,5881 khi thay feature ngẫu nhiên. Xác suất của lớp dự đoán ban đầu giảm lần lượt 31,23 và 14,64 điểm phần trăm. Theo lớp, SlowHTTPTest có win rate 82,25%, Infilteration 91,07% và FTP BruteForce 92,30%, thấp hơn mức tổng thể.")
    add_figure(
        doc,
        faithfulness_chart,
        "Hình 4. Mức giảm raw margin khi thay Top k feature SHAP và feature ngẫu nhiên. Nguồn: Báo cáo kết quả nghiên cứu SHAP trên XGBoost.",
        width=5.8,
    )
    add_body(doc, "Kết luận Sub RQ 2.3: Top 5 SHAP thắng đối chứng ở 695.140 trên 699.996 flow đủ điều kiện, tương đương 99,31%. Đây là bằng chứng về ảnh hưởng tương đối trong phép thử đã chọn. Không gọi 99,31% là accuracy của SHAP, không suy ra model phân loại đúng 99,31% và không xem đây là bằng chứng nhân quả ngoài model.")
    add_body(doc, "Kiểm tra additivity có sai số lớn nhất khoảng 2,2888 x 10 mũ trừ 5 raw margin. Nó chỉ chứng minh tổng đóng góp SHAP khớp gần với score model, không thay cho faithfulness.")

    add_heading(doc, "7 SHAP đáp ứng Sub RQ 2.1 đến đâu")
    add_body(doc, "SHAP đã có các bảng ranking và Top k, nhưng hai báo cáo chưa có bảng feature LIME cho cùng sample và cùng output class. Vì vậy, hiện chưa có tỷ lệ đồng thuận SHAP LIME và không được kết luận hai phương pháp giống nhau hoặc phương pháp nào tốt hơn.")
    add_body(doc, "Các chỉ số overlap, Jaccard và Spearman đã dùng trong báo cáo SHAP để kiểm tra độ ổn định giữa các lần lấy mẫu SHAP hoặc trước sau nhiễu. Chúng không phải similarity SHAP LIME.")

    add_heading(doc, "8 LIME cần làm gì tiếp theo")
    add_table(
        doc,
        ["Bước", "Yêu cầu lấy từ thiết kế RQ2"],
        [
            ["1", "Dùng cùng flow x và cùng output class c mà SHAP giải thích."],
            ["2", "Thống nhất k, cách xử lý hòa hạng và cách đổi điều kiện rời rạc LIME về tên feature gốc."],
            ["3", "Không ép đủ k bằng các hệ số bằng không."],
            ["4", "Nếu LIME giải thích xác suất còn SHAP giải thích raw margin, cần điều chỉnh thiết kế hoặc nêu rõ giới hạn so sánh."],
            ["5", "Tính Overlap at k bằng 100 x số feature chung chia k."],
            ["6", "Tính Jaccard at k bằng số feature chung chia số feature thuộc hợp của hai Top k."],
            ["7", "Tổng hợp similarity theo từng attack; sau đó đối chiếu Top feature LIME với kiến thức miền."],
            ["8", "Nếu muốn kết luận faithfulness cho cả hai phương pháp, áp dụng phép thay feature tương tự cho LIME."],
        ],
        [Inches(0.6), Inches(5.75)],
    )
    add_body(doc, "Sau LIME, RQ2 sẽ trả lời đầy đủ hơn: Sub RQ 2.1 có phần trăm similarity; Sub RQ 2.2 có đối chiếu feature của cả SHAP và LIME; Sub RQ 2.3 có faithfulness của từng phương pháp theo một thiết kế can thiệp rõ ràng.")

    add_heading(doc, "9 Kết luận tiến độ")
    add_body(doc, "SHAP đã hoàn thành phần việc độc lập cần thiết cho RQ2: giải thích feature theo lớp, tạo biểu đồ Top 5, phân tích cảnh báo cụ thể, kiểm tra liên hệ domain ở mức định tính và đo faithfulness có đối chứng. LIME là phần bắt buộc còn lại để tạo kết quả similarity SHAP LIME; sau đó cần đánh giá faithfulness LIME nếu RQ2 muốn so sánh độ trung thực của cả hai phương pháp.")
    add_body(doc, "Nguồn nội dung và hình trong báo cáo này chỉ gồm Bao_cao_RQ_SHAP_XGBoost.docx và Bao_cao_Top5_SHAP_tung_loai_tan_cong.docx.")

    doc.core_properties.title = "Báo cáo tiến độ SHAP cho RQ2"
    doc.core_properties.subject = "Mục đích RQ2 kết quả SHAP và công việc LIME tiếp theo"
    doc.core_properties.author = "Nhóm nghiên cứu"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    create_full_source_only_report()
