from docx import Document
import csv
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

# Script resides in RQ2_Tools; generated documents belong in RQ2_Tai_lieu.
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "RQ2_Tai_lieu" / "Huong_dan_chi_tiet_cac_chi_so_LIME_RQ2.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.8)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.14
for name, size in [('Title', 21), ('Heading 1', 15), ('Heading 2', 12)]:
    s = styles[name]
    s.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    s._element.rPr.rFonts.set(qn('w:ascii'), s.font.name)
    s._element.rPr.rFonts.set(qn('w:hAnsi'), s.font.name)
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0,0,0)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def border(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'6'); el.set(qn('w:color'),color)

def set_cell_text(cell, value, bold=False, white=False, center=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    r = p.add_run(str(value)); r.bold = bold; r.font.size = Pt(9.2)
    if white: r.font.color.rgb = RGBColor(255,255,255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, bold=True, white=True, center=True)
        shade(t.rows[0].cells[i], '17365D')
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, center=(i == 0 and len(row) <= 3))
            if ri % 2 == 1: shade(cells[i], 'F3F7FB')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def heading(text, level=1):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p

def p(text='', bold_lead=None):
    par = doc.add_paragraph()
    if bold_lead:
        r=par.add_run(bold_lead); r.bold=True
    par.add_run(text)
    return par

def bullets(items):
    for item in items:
        par = doc.add_paragraph(style='List Bullet')
        par.add_run(item)

# Title page / opening
title = doc.add_paragraph('Huong dan doc va dien giai ket qua LIME RQ2', style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph('Tai lieu rieng cho phan giai thich XAI LIME trong he thong phat hien xam nhap CICIDS2018')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].italic = True; sub.runs[0].font.size = Pt(11)
doc.add_paragraph()
p('Tài liệu này giải thích LIME và toàn bộ chỉ số đã được tính trong thư mục RQ2_LIME_FI. Kết luận cần nhớ: LIME đã chạy đúng quy trình ghép cặp với SHAP trên 1.339 flow dự đoán đúng, nhưng mô hình tuyến tính cục bộ của LIME có Local R² thấp. Vì vậy, kết quả LIME hữu ích để phân tích và so sánh, nhưng không nên dùng như bằng chứng duy nhất để khẳng định nguyên nhân dự đoán của XGBoost.')

heading('1 Boi canh thuc nghiem')
p('Mục tiêu của LIME trong RQ2 không phải đo độ chính xác của XGBoost. Độ chính xác đã thuộc RQ1. LIME trả lời câu hỏi: với một flow cụ thể mà XGBoost đã dự đoán, những đặc trưng mạng nào đang có ảnh hưởng cục bộ đến dự đoán đó?')
table(['Thành phần', 'Thiết lập đã dùng', 'Ý nghĩa'], [
['Mô hình được giải thích', 'XGBoost tốt nhất', 'Mô hình phân loại 15 nhãn đã được khóa từ RQ1.'],
['Tập flow giải thích', '1.339 flow dự đoán đúng', 'Lấy phân tầng, tối đa 100 flow mỗi lớp; SQL Injection có 9 flow và XSS có 44 flow.'],
['Đầu ra được giải thích', 'Lớp mà XGBoost dự đoán', 'Không giải thích nhãn thật. Vì cohort chỉ giữ flow đúng nên nhãn dự đoán trùng nhãn thật.'],
['Số đặc trưng', '78 feature CICFlowMeter', 'Giữ đúng thứ tự feature của model.'],
['LIME background', '8.002 flow', 'Dùng để tạo các mẫu lân cận và huấn luyện surrogate cục bộ.'],
['Lấy mẫu LIME', '5.000 mẫu mỗi flow', 'Số điểm lân cận được LIME tạo ra để xấp xỉ hành vi cục bộ của XGBoost.'],
['Kernel width', '6,624', 'Quy định mẫu ở gần flow gốc được ưu tiên mạnh hơn mẫu ở xa.'],
['Discretize continuous', 'Có', 'LIME chia feature liên tục thành khoảng để biểu diễn điều kiện cục bộ dễ đọc.'],
], [1.25, 1.65, 4.0])

heading('2 LIME la gi va no hoat dong nhu the nao')
p('LIME là viết tắt của Local Interpretable Model agnostic Explanations. “Model agnostic” nghĩa là LIME không cần biết cấu trúc bên trong của XGBoost; nó chỉ cần gọi được hàm dự đoán xác suất. “Local” nghĩa là kết luận chỉ dành cho một flow và vùng lân cận của flow đó, không phải quy luật chung cho toàn bộ dữ liệu.')
heading('2 1 Quy trinh cua mot loi giai thich LIME', 2)
table(['Bước', 'LIME thực hiện gì', 'Cách hiểu đơn giản'], [
['1', 'Chọn một flow gốc x', 'Ví dụ một flow được XGBoost dự đoán là SSH Brute Force.'],
['2', 'Sinh 5.000 flow lân cận', 'Thay đổi nhẹ các feature quanh flow gốc để xem xác suất dự đoán thay đổi ra sao.'],
['3', 'Gọi XGBoost cho từng flow lân cận', 'XGBoost cung cấp xác suất của lớp đang được giải thích.'],
['4', 'Gán trọng số theo khoảng cách', 'Flow càng gần x càng quan trọng; flow xa ảnh hưởng ít hơn.'],
['5', 'Fit Ridge regression cục bộ', 'LIME dùng một mô hình tuyến tính đơn giản để bắt chước XGBoost trong vùng đó.'],
['6', 'Xuất hệ số feature', 'Hệ số dương đẩy dự đoán về lớp đang xét; hệ số âm kéo dự đoán ra xa lớp đó.'],
], [0.45, 2.55, 3.9])
p('Có thể viết ý tưởng của LIME như sau: LIME tìm một mô hình đơn giản g sao cho g gần với dự đoán của XGBoost f ở quanh flow x, đồng thời g phải dễ hiểu. Chất lượng của “gần” này được đo bằng Local R². Vì XGBoost là phi tuyến, một đường thẳng cục bộ có thể không bắt chước tốt được nó.')

heading('3 Cach doc mot loi giai thich LIME')
p('Mỗi flow có danh sách feature và trọng số. Ví dụ “Dst Port có trọng số dương” không có nghĩa cổng đích luôn là nguyên nhân của mọi SSH Brute Force. Nó chỉ nói: với flow này, trong xấp xỉ cục bộ của LIME, giá trị Dst Port hiện tại góp phần làm xác suất lớp được giải thích tăng lên.')
table(['Khái niệm', 'Diễn giải đúng', 'Không được suy diễn thành'], [
['Top 5 feature', 'Năm feature có trị tuyệt đối hệ số LIME lớn nhất của một flow.', 'Năm feature quan trọng nhất trên toàn bộ dataset.'],
['Trọng số dương', 'Theo surrogate LIME, feature làm tăng điểm/xác suất lớp đang giải thích trong vùng cục bộ.', 'Feature chắc chắn gây ra tấn công.'],
['Trọng số âm', 'Theo surrogate, feature làm giảm điểm/xác suất lớp đang giải thích.', 'Feature là dấu hiệu benign tuyệt đối.'],
['Trung bình theo lớp', 'Tổng hợp các explanation của các flow trong cùng lớp.', 'Quy tắc nhân quả của lớp tấn công.'],
], [1.45, 3.25, 2.2])

heading('4 Chi so quan trong nhat Local R2')
p('Local R² là chỉ số quan trọng nhất khi đánh giá riêng LIME. Nó đo mô hình Ridge tuyến tính của LIME bắt chước xác suất XGBoost tốt đến mức nào trên các mẫu lân cận mà LIME vừa sinh ra.')
table(['Giá trị Local R²', 'Cách đọc thực tế'], [
['Gần 1', 'Surrogate tuyến tính tái tạo tốt hành vi cục bộ của XGBoost. Ranking feature có cơ sở đáng tin hơn.'],
['Khoảng 0,5 đến 0,7', 'Có thể dùng để diễn giải cẩn trọng; vẫn nên kiểm tra thêm stability và domain validation.'],
['Dưới 0,3', 'Surrogate bắt chước XGBoost kém. Top feature chỉ là mô tả của một xấp xỉ yếu, không nên kết luận mạnh.'],
['Âm', 'Surrogate còn kém hơn cách đoán bằng trung bình; explanation không đáng tin.'],
], [1.7, 5.2])
p('Kết quả của bạn: median Local R² = 0,0677 và mean Local R² = 0,1192 trên 1.339 flow. Ngưỡng kiểm tra là 0,30. Như vậy phần lớn surrogate cục bộ không mô tả tốt ranh giới phi tuyến của XGBoost. Đây là phát hiện phương pháp luận, không phải lỗi chạy code.')
p('Đã có kiểm tra chẩn đoán trên 60 flow: tăng số mẫu từ 5.000 lên 10.000 hoặc tắt discretization không làm R² cải thiện đáng kể. Vì vậy câu viết an toàn là: “Trong các cấu hình LIME đã thử, chưa tìm thấy cấu hình cải thiện đáng kể Local R².” Không viết “LIME sai hoàn toàn” và cũng không viết “LIME giải thích chính xác mô hình”.')

heading('5 Ket qua cai dat va feature hang dau')
p('Bảng Top 1 theo lớp cho thấy feature mà LIME thường gán trọng số lớn nhất trong tập flow đã giải thích. Nó chỉ dùng để quan sát xu hướng, vì Local R² thấp.')
table(['Lớp', 'Feature Top 1 theo LIME', 'Cách ghi trong luận văn'], [
['Benign', 'FIN Flag Cnt', 'Xu hướng cục bộ, cần diễn giải thận trọng.'],
['Bot', 'Dst Port', 'Cổng đích xuất hiện nổi bật trong surrogate.'],
['Brute Force Web', 'RST Flag Cnt', 'Không khẳng định đây là chữ ký tấn công.'],
['Brute Force XSS', 'Bwd Pkts/s', 'Thông tin flow, không phải nội dung payload XSS.'],
['DDoS HOIC', 'Init Fwd Win Byts', 'Xu hướng riêng của cohort đã chọn.'],
['DDoS LOIC UDP', 'Tot Fwd Pkts', 'Phù hợp trực giác lưu lượng cao, nhưng vẫn là LIME cục bộ.'],
['DDoS attacks HTTP', 'Fwd Pkts/s', 'Feature tốc độ gói theo chiều đi.'],
['DoS GoldenEye', 'Fwd Seg Size Min', 'Feature kích thước segment.'],
['DoS Hulk', 'FIN Flag Cnt', 'Cần đối chiếu thêm với SHAP/domain.'],
['DoS SlowHTTPTest', 'Fwd Seg Size Min', 'Không là bằng chứng payload HTTP.'],
['DoS slowloris', 'FIN Flag Cnt', 'Chỉ mô tả output LIME của cohort.'],
['FTP BruteForce', 'Fwd Seg Size Min', 'Không dùng một feature để kết luận lớp.'],
['Infilteration', 'Fwd URG Flags', 'Cần kiểm tra domain và số mẫu.'],
['SQL Injection', 'Bwd IAT Std', 'Chỉ có 9 flow nên độ ổn định theo lớp hạn chế.'],
['SSH BruteForce', 'Dst Port', 'Có trực giác dịch vụ SSH, nhưng không chứng minh hành vi thử mật khẩu.'],
], [1.55, 2.0, 3.35])

heading('6 Similarity giua LIME va SHAP')
p('Phần này thuộc Sub RQ 2.1. Cả hai phương pháp được chạy trên cùng 1.339 flow và cùng lớp dự đoán. SHAP đo đóng góp vào raw margin của XGBoost, còn LIME fit surrogate theo xác suất. Do hai thang đo khác nhau, chỉ so sánh danh sách feature, thứ hạng và chiều dấu; không so trực tiếp độ lớn SHAP với độ lớn hệ số LIME.')
table(['Chỉ số', 'Kết quả macro', 'Ý nghĩa'], [
['Jaccard at 3', '0,241', 'Mức giao nhau Top 3 feature thấp. Công thức: số feature chung chia số feature thuộc hợp của hai Top k.'],
['Jaccard at 5', '0,213', 'Top 5 của LIME và SHAP chỉ trùng khoảng 21,3 phần trăm theo trung bình lớp.'],
['Jaccard at 10', '0,239', 'Mở rộng Top 10 không làm mức trùng lặp tăng nhiều.'],
['Spearman at 3', '0,818', 'Khi feature đã cùng xuất hiện, thứ hạng ở Top 3 tương đối giống nhau.'],
['Spearman at 5', '0,476', 'Tương quan thứ hạng mức trung bình.'],
['Spearman at 10', '0,380', 'Thứ hạng càng sâu càng khác nhau.'],
['Signed agreement at 5', '0,896', 'Các feature trùng nhau thường có cùng chiều dương hoặc âm. Đây chỉ là chỉ số phụ vì đầu ra hai phương pháp khác thang đo.'],
], [1.55, 1.25, 4.1])
p('Kết luận đúng: LIME và SHAP thường không chọn cùng Top 5 feature, nhưng các feature hiếm khi trùng thì đa số có cùng chiều tác động. Không kết luận rằng hai phương pháp “tương đương”.')

heading('7 Domain Precision at 5')
p('Domain Precision at 5 là phép kiểm tra kiến thức miền: Top 5 feature của một explanation có bao nhiêu feature nằm trong danh sách tín hiệu được kỳ vọng trước cho loại tấn công đó. Công thức: số feature Top 5 khớp expected signals chia 5. Danh sách expected signals được thiết kế dựa trên cơ chế tấn công, không được chọn sau khi xem kết quả.')
table(['Phương pháp', 'Macro Domain Precision at 5', 'Diễn giải'], [
['SHAP', '0,187', 'Trung bình khoảng 18,7 phần trăm feature Top 5 khớp danh sách expected signals.'],
['LIME', '0,111', 'Trung bình khoảng 11,1 phần trăm feature Top 5 khớp danh sách expected signals.'],
], [1.6, 2.2, 3.1])
p('Hai giá trị đều thấp hơn 20 phần trăm. Điều này không chứng minh model sai: CICFlowMeter là feature thống kê flow, trong khi hiểu biết domain có thể mong đợi payload, chuỗi HTTP, tên tiến trình hoặc ngữ cảnh mà dataset không chứa. Kết quả cho thấy Top 5 giải thích chưa khớp mạnh với bộ tín hiệu kỳ vọng đã định nghĩa, và LIME thấp hơn SHAP trong phép đo này.')

heading('8 Stability va Robustness')
heading('8 1 Stability qua cac lan chay LIME', 2)
p('LIME có ngẫu nhiên nội tại vì nó phải sinh các mẫu lân cận. Stability hỏi: cùng một flow, nếu đổi seed và chạy lại LIME, Top 5 feature có giữ nguyên không? Thí nghiệm dùng 289 flow, tối đa 20 flow mỗi lớp, 8 seed và 5.000 mẫu mỗi lần chạy.')
table(['Chỉ số', 'Kết quả', 'Cách hiểu'], [
['Median Jaccard Top 5', '0,415', 'Một nửa các so sánh có mức giao Top 5 không quá khoảng 41,5 phần trăm. Ổn định ở mức trung bình, không cao.'],
['Phạm vi', '289 flow', 'Không suy rộng con số này thành toàn bộ 1.339 flow hay toàn bộ test set.'],
], [1.8, 1.25, 3.85])
heading('8 2 Robustness khi nhieu input', 2)
p('Robustness hỏi: nếu làm nhiễu nhẹ feature thời gian/tốc độ của flow, dự đoán và Top 5 LIME thay đổi thế nào? Đây không phải tấn công thật; nó là stress test độ nhạy quanh input.')
table(['Mức nhiễu', 'Mean Jaccard Top 5', 'Median Jaccard', 'Tỷ lệ đổi dự đoán'], [
['1 phần trăm', '0,450', '0,429', '3,00 phần trăm'],
['5 phần trăm', '0,449', '0,429', '4,15 phần trăm'],
['10 phần trăm', '0,448', '0,429', '5,77 phần trăm'],
], [1.1, 1.55, 1.45, 1.65])
p('Dù phần lớn nhãn dự đoán không đổi, Top 5 explanation chỉ giao nhau xấp xỉ 45 phần trăm. Điều này phù hợp với phát hiện Local R² thấp và stability trung bình: explanation LIME nhạy hơn nhãn model.')

heading('9 Faithfulness bang donor substitution')
p('Faithfulness kiểm tra trực tiếp hơn: nếu chọn feature mà explanation nói là quan trọng và thay chúng bằng giá trị lấy từ các flow train khác, xác suất của lớp dự đoán ban đầu có giảm mạnh hơn so với thay ngẫu nhiên hay không? Đây không kiểm tra LIME bắt chước XGBoost tốt như Local R²; nó là phép thử khác về mức ảnh hưởng của feature được chọn.')
table(['Thiết kế chung cho SHAP và LIME', 'Giá trị'], [
['Cohort', 'Cùng 1.339 flow, cùng lớp dự đoán ban đầu.'],
['Donor', '10 dòng train, seed 173.'],
['Đối chứng', 'Chọn ngẫu nhiên k feature, 20 lần lặp, seed 2026.'],
['Top k', 'k = 3, 5, 10. Chọn feature đóng góp dương.'],
['Thắng', 'Độ giảm xác suất do Top k lớn hơn random control quá 10 mũ trừ 6.'],
], [2.45, 4.05])
table(['k', 'LIME win rate', 'LIME mean probability drop', 'SHAP win rate', 'Ý nghĩa'], [
['3', '82,60 phần trăm', '7,69 điểm phần trăm', '98,88 phần trăm', 'SHAP chọn feature gây giảm xác suất mạnh hơn random thường xuyên hơn.'],
['5', '90,44 phần trăm', '8,07 điểm phần trăm', '95,74 phần trăm', 'LIME khá hơn khi cho phép chọn 5 feature nhưng vẫn thấp hơn SHAP.'],
['10', '94,77 phần trăm', '8,81 điểm phần trăm', '99,81 phần trăm', 'SHAP chỉ có coverage 80,36 phần trăm ở k 10; so sánh phải nêu coverage.'],
], [0.45, 1.25, 1.85, 1.2, 2.0])
p('Không được lấy Faithfulness cao để bỏ qua Local R² thấp. Hai chỉ số trả lời hai câu hỏi khác nhau: Local R² hỏi surrogate LIME có bắt chước XGBoost không; donor substitution hỏi feature được chọn có tác động khi bị thay thế không. Trong báo cáo của bạn, LIME có faithfulness tương đối tốt ở k lớn nhưng chất lượng surrogate vẫn là hạn chế chính.')

heading('10 Cach viet ket luan LIME trong RQ2')
p('Đoạn có thể dùng trong luận văn: “LIME được triển khai trên cohort 1.339 flow dự đoán đúng, sử dụng 78 đặc trưng CICFlowMeter và giải thích lớp do XGBoost dự đoán. Kết quả cho thấy Local R² trung vị 0,0677, thấp hơn ngưỡng 0,30, cho thấy surrogate tuyến tính cục bộ chưa xấp xỉ tốt hành vi phi tuyến của XGBoost trong các cấu hình đã thử. Vì vậy các Top feature LIME được dùng như bằng chứng mô tả và so sánh, không phải bằng chứng độc lập về nguyên nhân dự đoán. Khi so sánh ghép cặp với SHAP, Jaccard at 5 đạt 0,213, Domain Precision at 5 đạt 0,111 và faithfulness của LIME thấp hơn SHAP ở k bằng 3, 5 và 10. Do đó SHAP là phương pháp phù hợp hơn để làm kết luận chính của RQ2, còn LIME đóng vai trò đối chứng cục bộ và chỉ ra giới hạn của surrogate tuyến tính.”')

heading('11 Nhung dieu can tranh khi bao cao')
bullets([
'Không nói Local R² là độ chính xác phân loại của XGBoost.',
'Không nói một Top feature LIME là nguyên nhân gây tấn công.',
'Không so trực tiếp trị tuyệt đối hệ số LIME với trị tuyệt đối SHAP vì khác thang đo đầu ra.',
'Không suy rộng kết quả stability và robustness từ 289 flow lên toàn bộ dataset.',
'Không diễn giải quá mạnh SQL Injection và XSS vì số flow trong cohort nhỏ, đặc biệt SQL Injection chỉ có 9 flow.',
'Không gộp kết quả faithfulness cohort 1.339 flow với kết quả SHAP chạy trên test set lớn ở quy trình khác.',
])

doc.add_page_break()
heading('Phu luc A Top 5 LIME day du theo tung lop')
p('Bảng dưới đây lấy trực tiếp từ lime_top5_by_class.csv. Đây là Top 5 trung bình theo lớp của cohort ghép cặp. Cột “trọng số tuyệt đối trung bình” dùng để xếp hạng mức độ nổi bật của feature trong các explanation LIME của lớp đó. Nó không cho biết chiều tác động dương hay âm và không phải là feature quan trọng toàn cục của XGBoost.')
rows = []
with open(r"C:\Users\LOQ\Downloads\SELF_KLTN\RQ2_LIME_FI\lime_top5_by_class.csv", encoding='utf-8-sig', newline='') as f:
    for r in csv.DictReader(f):
        rows.append([r['true_label'], str(sum(1 for old in rows if old[0] == r['true_label']) + 1), r['feature'], f"{float(r['abs_weight']):.6f}"])
table(['Lớp', 'Hạng', 'Feature', 'Trọng số tuyệt đối trung bình'], rows, [1.55, 0.55, 2.5, 1.8])

heading('12 Tom tat de nho nhanh')
table(['Câu hỏi', 'Trả lời từ kết quả hiện tại'], [
['LIME đã chạy được chưa?', 'Đã chạy đầy đủ các phần: per flow, agreement, domain, stability robustness và faithfulness.'],
['Có thể tin Top feature LIME như kết luận chính không?', 'Không nên, vì Local R² trung vị 0,0677 là thấp.'],
['LIME có hoàn toàn vô giá trị không?', 'Không. Nó vẫn là đối chứng cục bộ, cho thấy độ ổn định trung bình và cho phép so sánh công bằng với SHAP.'],
['Phương pháp nào nên làm kết luận chính RQ2?', 'SHAP, vì trong cohort ghép cặp nó có Domain Precision và Faithfulness cao hơn LIME, đồng thời TreeExplainer phù hợp với cấu trúc cây.'],
['Thông điệp học thuật chính?', 'XAI cần được đánh giá bằng nhiều chỉ số; chỉ có Top feature là chưa đủ.'],
], [2.45, 4.05])

# footer
for section in doc.sections:
    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = foot.add_run('RQ2 LIME | CICIDS2018 | Tai lieu dien giai ket qua')
    rr.font.size = Pt(8); rr.font.color.rgb = RGBColor(100,100,100)

doc.core_properties.title = 'Huong dan doc va dien giai ket qua LIME RQ2'
doc.core_properties.subject = 'Giai thich cac chi so LIME cho RQ2'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
