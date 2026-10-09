"""Tiện ích chung để dựng 4 báo cáo .docx của RQ2 bằng python-docx."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def new_doc(title: str, subtitle: str = "") -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Calibri")

    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(11)
    return doc


def add_heading(doc: Document, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False, italic: bool = False, size: int | None = None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table_from_df(doc: Document, df: pd.DataFrame, max_rows: int | None = 40, float_fmt: str = "{:.4f}"):
    df_show = df.copy()
    if max_rows is not None and len(df_show) > max_rows:
        df_show = df_show.head(max_rows)
    table = doc.add_table(rows=1, cols=len(df_show.columns))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for j, col in enumerate(df_show.columns):
        hdr[j].text = str(col)
        for p in hdr[j].paragraphs:
            for r in p.runs:
                r.bold = True
    for _, row in df_show.iterrows():
        cells = table.add_row().cells
        for j, val in enumerate(row):
            if isinstance(val, float):
                cells[j].text = float_fmt.format(val)
            else:
                cells[j].text = str(val)
    return table


def add_image(doc: Document, path: Path, width_cm: float = 14.0, caption: str | None = None):
    if not Path(path).exists():
        add_para(doc, f"[Thiếu hình: {path}]", italic=True)
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    last_p = doc.paragraphs[-1]
    last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.italic = True
        run.font.size = Pt(10)


def add_page_break(doc: Document):
    doc.add_page_break()


def add_image_with_analysis(
    doc: Document,
    path: Path,
    caption: str,
    analysis: list[str],
    width_cm: float = 14.0,
):
    """Chèn một hình + chú thích ngắn + phần PHÂN TÍCH đầy đủ ngay dưới hình (bullet list).

    Dùng cho mọi biểu đồ trong báo cáo tổng hợp — mỗi hình PHẢI đi kèm phân tích, không
    chỉ là ảnh trang trí."""
    add_image(doc, path, width_cm=width_cm, caption=caption)
    cap = doc.add_paragraph()
    run = cap.add_run("Phân tích hình trên:")
    run.bold = True
    add_bullets(doc, analysis)
