"""Build the RQ2 SHAP progress report from generated CSV outputs."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "RQ2_SHAP" / "outputs"
DOCX = ROOT / "RQ2_SHAP" / "Bao_Cao_Tien_Do_RQ2_SHAP.docx"
BLUE = "17365D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_table(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER)


def set_run(run, size=10.5, bold=False, color=None) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell(cell, value, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    set_run(run, 9.2, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    border_table(table)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, BLUE)
        set_cell(cell, header, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2 == 1:
                shade(cells[index], LIGHT_BLUE)
            if widths:
                cells[index].width = Inches(widths[index])
            set_cell(cells[index], value, align=WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def paragraph(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.14
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run(r, 10.5, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run(r, 10.5)
    else:
        r = p.add_run(text)
        set_run(r, 10.5)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run(r, 14 if level == 1 else 11.5, bold=True, color="000000")


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run(r, 10.3)


def main() -> None:
    manifest = pd.read_csv(OUT / "rq2_shared_samples.csv")
    prediction = pd.read_csv(OUT / "rq2_shap_sample_predictions.csv")
    ranking = pd.read_csv(OUT / "rq2_shap_ranking_all_features.csv")
    additivity = pd.read_csv(OUT / "rq2_shap_additivity_check.csv")

    total = len(manifest)
    correct = int(prediction["is_correct"].sum())
    accuracy = correct / total * 100
    sample_counts = manifest.groupby("true_label").size().sort_index()
    class_accuracy = prediction.groupby("true_label")["is_correct"].agg(["sum", "count"])
    class_accuracy["percent"] = class_accuracy["sum"] / class_accuracy["count"] * 100
    top1 = (
        ranking[ranking["rank"] == 1]
        .groupby("true_label")["feature"]
        .value_counts()
        .groupby(level=0)
        .head(1)
    )

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("BÁO CÁO TIẾN ĐỘ RQ2 SHAP")
    set_run(r, 18, bold=True, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Top k đặc trưng, lý do chọn 1.363 flow và kế hoạch LIME")
    set_run(r, 11, color="404040")
    doc.add_paragraph()

    heading(doc, "1. Mục tiêu và trạng thái hiện tại")
    paragraph(doc, "RQ2 đánh giá lời giải thích của SHAP và LIME cho các quyết định phân loại của XGBoost. Phần SHAP đã hoàn thành việc tạo tập flow chung, xếp hạng 78 đặc trưng của từng flow và xuất Top-1, Top-3, Top-5. Jaccard và Overlap chưa được tính vì cần LIME chạy trên đúng các flow này.")
    add_table(doc, ["Hạng mục", "Trạng thái", "Kết quả"], [
        ["Tập mẫu XAI", "Hoàn thành", f"{total:,} flow, seed 42"],
        ["SHAP cho từng flow", "Hoàn thành", "Xếp hạng đầy đủ 78 đặc trưng"],
        ["Top-k", "Hoàn thành", "Top-1, Top-3 và Top-5"],
        ["Kiểm tra cộng dồn SHAP", "Hoàn thành", f"Sai số TB {additivity.absolute_error.mean():.2e}; lớn nhất {additivity.absolute_error.max():.2e}"],
        ["LIME và so sánh", "Chưa thực hiện", "Phải dùng lại manifest SHAP"],
    ], [1.65, 1.3, 3.9])

    heading(doc, "2. Vì sao chọn 1.363 flow")
    paragraph(doc, "Tập test có 15 lớp và chênh lệch mạnh về số mẫu. Nếu lấy ngẫu nhiên toàn bộ test set, Benign và các lớp DDoS lớn sẽ chi phối kết quả, còn các lớp hiếm như SQL Injection gần như biến mất. Vì vậy, tập XAI được lấy mẫu phân tầng theo nhãn thật với seed 42.")
    paragraph(doc, "Mỗi lớp có ít nhất 100 mẫu được lấy đúng 100 flow. Hai lớp hiếm không đủ 100 mẫu được giữ toàn bộ: Brute Force-XSS có 46 flow và SQL Injection có 17 flow. Cách chọn này vừa bảo đảm mọi loại tấn công đều có mặt, vừa giới hạn chi phí LIME ở bước tiếp theo.")
    add_table(doc, ["Nhóm lớp", "Quy tắc chọn", "Số flow"], [
        ["13 lớp có từ 100 mẫu trở lên", "Lấy ngẫu nhiên phân tầng 100 flow/lớp", "1.300"],
        ["Brute Force-XSS", "Lấy toàn bộ vì chỉ có 46 flow test", "46"],
        ["SQL Injection", "Lấy toàn bộ vì chỉ có 17 flow test", "17"],
        ["Tổng", "Dùng chung cho SHAP và LIME", f"{total:,}"],
    ], [2.1, 3.55, 1.2])
    paragraph(doc, "sample_id là vị trí dòng gốc trong test set. Đây là khóa nối bắt buộc: LIME phải giải thích đúng sample_id mà SHAP đã giải thích. Nhờ đó, so sánh giữa hai phương pháp là công bằng trên cùng dữ liệu, cùng model và cùng quyết định dự đoán.")

    heading(doc, "3. Quy trình SHAP đã chạy")
    add_table(doc, ["Bước", "Thực hiện"], [
        ["1", "Nạp test_set.parquet, XGBoost và LabelEncoder đã lưu"],
        ["2", "Kiểm tra 78 feature và giữ nguyên thứ tự feature của model"],
        ["3", "Tạo manifest 1.363 flow với seed 42"],
        ["4", "Dùng TreeExplainer để tính SHAP cho mô hình XGBoost đa lớp"],
        ["5", "Với mỗi flow, giải thích lớp mà XGBoost dự đoán"],
        ["6", "Xếp hạng theo trị tuyệt đối SHAP và xuất Top-1, Top-3, Top-5"],
        ["7", "Kiểm tra base value cộng tổng SHAP khớp raw margin của XGBoost"],
    ], [0.65, 6.15])
    paragraph(doc, "Kết quả prediction trên tập XAI là " + f"{correct:,}/{total:,} flow đúng, tương đương {accuracy:.2f}%. Con số này mô tả XGBoost trên tập mẫu XAI, không phải độ chính xác của SHAP.")
    paragraph(doc, "SHAP giải thích điểm raw của lớp dự đoán. Sai số cộng dồn rất nhỏ cho thấy phép phân rã SHAP khớp về mặt số học với output model. Kiểm tra này không thay thế faithfulness hoặc xác nhận model dự đoán đúng nhãn thật.")

    heading(doc, "4. Tóm tắt Top 1 SHAP hiện tại")
    rows = []
    for label, count in sample_counts.items():
        result = class_accuracy.loc[label]
        feature, frequency = top1.loc[label].index[0], int(top1.loc[label].iloc[0])
        rows.append([label, int(count), f"{int(result['sum'])}/{int(result['count'])} ({result['percent']:.1f}%)", feature, frequency])
    add_table(doc, ["Lớp thật", "Mẫu", "Dự đoán đúng", "Top-1 SHAP xuất hiện nhiều nhất", "Lần"], rows, [1.55, 0.5, 1.1, 2.9, 0.4])
    paragraph(doc, "Bảng này chỉ là tổng hợp Top-1 theo nhãn thật. Với mẫu dự đoán sai, SHAP vẫn giải thích lớp mà XGBoost dự đoán, vì vậy cần đọc kết quả cùng cột predicted_label và is_correct trong các CSV chi tiết.")

    heading(doc, "5. Kế hoạch thực hiện LIME")
    paragraph(doc, "LIME không được lấy một tập mẫu mới. LIME phải đọc rq2_shared_samples.csv và giải thích cùng 1.363 flow, cùng 78 feature, cùng model XGBoost và cùng predicted_label đã lưu từ bước SHAP.")
    add_table(doc, ["Hạng mục", "Quy định cho LIME"], [
        ["Tập flow", "Dùng đúng rq2_shared_samples.csv"],
        ["Lớp giải thích", "Giải thích lớp XGBoost dự đoán cho từng sample_id"],
        ["Cấu hình", "num_samples = 1500; kernel_width = 0.75; num_features = 5"],
        ["Kết quả cần lưu", "sample_id, predicted_label, rank, feature, feature_key, lime_weight, local_r2"],
        ["Tên feature", "Chuẩn hóa về feature_key giống SHAP; tách điều kiện LIME như Feature > value về tên feature gốc"],
        ["Đầu ra", "Lưu full ranking nếu có thể, tối thiểu Top-1, Top-3, Top-5"],
    ], [1.7, 5.1])

    heading(doc, "6. So sánh SHAP và LIME để trả lời RQ2.1")
    paragraph(doc, "Sau khi có LIME, ghép hai kết quả theo sample_id. Với từng flow, tạo tập Top-k SHAP và Top-k LIME, sau đó tính k = 1, 3, 5. Top-5 là mức phân tích chính vì LIME được cấu hình hiển thị năm đặc trưng.")
    add_table(doc, ["Chỉ số", "Công thức", "Ý nghĩa"], [
        ["Top-1 agreement", "1 nếu cùng feature Top-1, ngược lại 0", "Hai phương pháp có đồng ý về lý do mạnh nhất không"],
        ["Overlap Top-k", "|SHAP ∩ LIME| / k × 100", "Tỷ lệ feature chung, dễ diễn giải"],
        ["Jaccard Top-k", "|SHAP ∩ LIME| / |SHAP ∪ LIME| × 100", "Độ tương đồng chặt hơn, tính cả feature bất đồng"],
    ], [1.45, 2.75, 2.6])
    paragraph(doc, "Ví dụ: nếu SHAP Top-5 và LIME Top-5 có 3 feature chung, Overlap là 3/5 = 60%, còn Jaccard là 3/7 = 42,86%. Báo cáo cần tính trung bình theo từng lớp thật, sau đó tính trung bình toàn bộ tập và đọc riêng nhóm dự đoán đúng/sai.")

    heading(doc, "7. Domain Validation để trả lời RQ2.2")
    paragraph(doc, "Sau khi có Top-5 của cả SHAP và LIME, đối chiếu feature với cơ chế kỹ thuật đã biết. Kết luận chỉ nên dùng ba mức: phù hợp, phù hợp một phần, hoặc chưa đủ bằng chứng.")
    add_table(doc, ["Nhóm tấn công", "Dấu hiệu flow kỳ vọng", "Giới hạn khi diễn giải"], [
        ["DDoS và DoS flood", "Packets/s, Bytes/s, số gói, Flow Duration", "Cần xem hướng đóng góp và kết hợp nhiều feature"],
        ["SlowHTTPTest và Slowloris", "Flow Duration, IAT, Bytes/s thấp", "Flow feature không mô tả đầy đủ HTTP header"],
        ["SSH và FTP brute force", "Dst Port, IAT, packet rate, số gói forward", "Port là dấu hiệu dịch vụ, không tự chứng minh brute force"],
        ["Bot và Infilteration", "Port, header, flags, timing", "Thường là dấu hiệu gián tiếp; cần log hoặc PCAP"],
        ["Web brute force, SQLi và XSS", "Kích thước gói, port, IAT", "Không quan sát trực tiếp payload hay nội dung request"],
    ], [1.7, 2.35, 2.75])

    heading(doc, "8. Faithfulness để trả lời RQ2.3")
    paragraph(doc, "Không dùng một tỷ lệ chung gọi là SHAP chính xác bao nhiêu phần trăm. Với SHAP, cần kiểm tra xem feature được chọn có thật sự ảnh hưởng đến output XGBoost hay không. Với mỗi flow, giữ cố định lớp dự đoán gốc, thay Top-3 và Top-5 feature SHAP dương bằng giá trị tham chiếu từ train set, sau đó chạy lại model.")
    add_table(doc, ["So sánh", "Cần đo"], [
        ["Top-k SHAP", "Mức giảm raw margin, giảm xác suất lớp gốc và tỷ lệ đổi nhãn"],
        ["Random baseline", "Thay ngẫu nhiên cùng k feature trong tập feature SHAP dương"],
        ["Kết luận", "Tỷ lệ flow mà Top-k SHAP gây giảm mạnh hơn random, kèm mức giảm trung bình"],
        ["LIME", "Làm cùng thiết kế để so sánh công bằng; thêm Local R² của LIME"],
    ], [1.65, 5.15])

    heading(doc, "9. Bộ kết quả cuối cùng cho RQ2")
    add_table(doc, ["Câu hỏi", "Bằng chứng cần nộp", "Trạng thái"], [
        ["RQ2.1 Similarity", "Overlap và Jaccard Top-1/3/5 theo lớp và toàn bộ tập", "Chờ LIME"],
        ["RQ2.2 Domain Validation", "Bảng SHAP và LIME đối chiếu với cơ chế tấn công", "SHAP sẵn sàng; chờ LIME"],
        ["RQ2.3 Faithfulness", "Top-3/5 SHAP và LIME so với random baseline", "Chưa chạy"],
        ["RQ3 Dashboard", "Dùng SHAP Top-3/5 cho giải thích cảnh báo", "Có thể dùng sau khi chốt RQ2"],
    ], [1.55, 4.25, 1.1])
    paragraph(doc, "Thứ tự tiếp theo: chạy LIME trên manifest hiện có, chuẩn hóa tên feature, tính Overlap và Jaccard, lập bảng domain validation chung, rồi thực hiện faithfulness cho SHAP và LIME bằng cùng baseline ngẫu nhiên.")

    heading(doc, "10. Tệp kết quả hiện có")
    for item in [
        "rq2_shared_samples.csv: danh sách 1.363 flow bắt buộc dùng lại cho LIME.",
        "rq2_shap_ranking_all_features.csv: đủ 78 feature cho từng flow.",
        "rq2_shap_top1.csv, rq2_shap_top3.csv, rq2_shap_top5.csv: dữ liệu Top-k theo sample_id.",
        "rq2_shap_top5_by_class.csv: Top-5 trung bình theo lớp thật và trạng thái dự đoán đúng/sai.",
        "rq2_shap_additivity_check.csv: bằng chứng kiểm tra cộng dồn SHAP.",
    ]:
        bullet(doc, item)

    doc.add_page_break()
    heading(doc, "11. Hướng dẫn đọc từng tệp kết quả")
    paragraph(doc, "Các tệp nằm trong thư mục RQ2_SHAP/outputs. Khi xem lần đầu, mở README_results.md trước vì đây là bảng tóm tắt dễ đọc nhất. Sau đó mở các CSV nhỏ theo thứ tự dưới đây, không cần mở ranking 78 feature ngay từ đầu.")
    add_table(doc, ["Tệp", "Dùng để làm gì", "Nên mở khi nào"], [
        ["README_results.md", "Tóm tắt số mẫu mỗi lớp, prediction và Top-1 SHAP thường gặp", "Mở đầu tiên"],
        ["rq2_shared_samples.csv", "Manifest 1.363 flow; là danh sách chuẩn để LIME chạy lại", "Trước khi làm LIME"],
        ["rq2_shap_sample_predictions.csv", "So sánh nhãn thật, nhãn XGBoost dự đoán, xác suất và đúng/sai", "Khi phân tích lỗi model"],
        ["rq2_shap_top1.csv", "Một feature quan trọng nhất cho mỗi flow; 1.363 dòng", "Xem lý do mạnh nhất"],
        ["rq2_shap_top3.csv", "Ba feature quan trọng nhất; 4.089 dòng", "Dùng cho diễn giải alert ngắn"],
        ["rq2_shap_top5.csv", "Năm feature quan trọng nhất; 6.815 dòng", "Dùng chính cho RQ2 và so với LIME"],
        ["rq2_shap_top5_by_class.csv", "Top-5 trung bình theo lớp thật và nhóm dự đoán đúng/sai", "Dùng cho Domain Validation"],
        ["rq2_shap_ranking_all_features.csv", "Đủ 78 feature cho từng flow; 106.314 dòng", "Dùng cho phân tích sâu hoặc đổi k sau này"],
        ["rq2_shap_additivity_check.csv", "Đối chiếu raw margin và tổng SHAP", "Dùng làm bằng chứng kiểm tra kỹ thuật"],
        ["rq2_feature_schema.json", "Danh sách 78 feature và feature_key chuẩn hóa", "Dùng để map tên LIME về đúng feature gốc"],
    ], [2.2, 3.0, 1.35])

    heading(doc, "12. Ý nghĩa các cột quan trọng")
    add_table(doc, ["Cột", "Ý nghĩa và cách đọc"], [
        ["sample_id", "Vị trí dòng gốc, đánh số từ 0, trong test_set.parquet. Đây là khóa để SHAP và LIME giải thích đúng cùng một flow."],
        ["true_label", "Nhãn thật trong dữ liệu CICIDS2018. Dùng để nhóm kết quả theo từng loại tấn công."],
        ["predicted_label", "Lớp mà XGBoost thực sự dự đoán. SHAP giải thích lớp này, kể cả khi dự đoán sai."],
        ["predicted_probability", "Xác suất do model gán cho lớp dự đoán. Dùng để tham khảo độ mạnh output model, không gọi là xác suất đã hiệu chuẩn."],
        ["is_correct", "True nếu predicted_label trùng true_label; False nếu XGBoost dự đoán sai."],
        ["rank", "Thứ hạng feature trong một flow. rank = 1 là quan trọng nhất theo trị tuyệt đối SHAP; rank = 78 là thấp nhất."],
        ["feature", "Tên feature đúng như lúc XGBoost train, ví dụ Flow Pkts/s hoặc Fwd IAT Mean."],
        ["feature_key", "Tên chuẩn hóa viết thường, bỏ khoảng trắng và dấu câu. Cột này dùng để ghép SHAP với LIME sau này."],
        ["shap_value", "Đóng góp có dấu cho predicted_label. Giá trị dương đẩy model về lớp dự đoán; giá trị âm kéo model ra khỏi lớp đó."],
        ["abs_shap_value", "Trị tuyệt đối của SHAP value. Đây là cơ sở xếp hạng Top-k vì đo độ mạnh đóng góp, không xét chiều."],
    ], [1.75, 4.8])

    heading(doc, "13. Top 1 Top 3 và Top 5 là gì")
    paragraph(doc, "Một flow có 78 feature, nhưng không cần hiển thị cả 78 feature cho người đọc. Top-k là k feature có abs_shap_value lớn nhất của flow đó. Trong bộ kết quả hiện có, Top-1 chứa rank 1; Top-3 chứa rank 1 đến 3; Top-5 chứa rank 1 đến 5.")
    add_table(doc, ["Tệp", "Số dòng", "Ý nghĩa nghiên cứu"], [
        ["rq2_shap_top1.csv", f"{total:,}", "Kiểm tra SHAP và LIME có đồng ý về lý do mạnh nhất không"],
        ["rq2_shap_top3.csv", f"{total * 3:,}", "Diễn giải ngắn cho một alert; dùng thêm trong faithfulness Top-3"],
        ["rq2_shap_top5.csv", f"{total * 5:,}", "Mức phân tích chính cho RQ2.1, Domain Validation và LIME"],
    ], [2.2, 1.2, 3.15])
    paragraph(doc, "Ví dụ, nếu một sample_id xuất hiện năm lần trong rq2_shap_top5.csv, năm dòng đó là năm feature quan trọng nhất của đúng flow đó. Muốn xem riêng một flow, lọc cột sample_id; sau này thực hiện y hệt trên file LIME, rồi so sánh hai tập feature.")

    heading(doc, "14. Cách chạy lại trong VS Code")
    paragraph(doc, "Mở thư mục gốc SELF_KLTN bằng Visual Studio Code, sau đó mở Terminal. Python 3.14 trên máy đã được cài các thư viện cần thiết cho lần chạy hiện tại: numpy, pandas, pyarrow, scikit-learn, xgboost, shap và joblib.")
    add_table(doc, ["Mục đích", "Lệnh trong VS Code Terminal"], [
        ["Đi đến thư mục dự án", "cd C:\\Users\\LOQ\\Downloads\\SELF_KLTN"],
        ["Chỉ kiểm tra dữ liệu và tạo lại manifest", "py -3.14 RQ2_SHAP\\run_rq2_shap_topk.py --dry-run"],
        ["Chạy lại SHAP đầy đủ", "py -3.14 RQ2_SHAP\\run_rq2_shap_topk.py"],
        ["Chạy với số mẫu nhỏ để thử", "py -3.14 RQ2_SHAP\\run_rq2_shap_topk.py --max-per-class 10"],
        ["Đổi seed nhưng giữ cách lấy mẫu", "py -3.14 RQ2_SHAP\\run_rq2_shap_topk.py --seed 123"],
    ], [2.15, 4.4])
    paragraph(doc, "Lệnh chạy đầy đủ sẽ ghi đè các CSV trong outputs bằng kết quả mới. Nếu muốn giữ kết quả seed 42 hiện tại, hãy sao chép thư mục outputs trước khi đổi seed hoặc đổi max-per-class.")

    heading(doc, "15. Cách xem kết quả trong VS Code")
    paragraph(doc, "Trong Explorer của VS Code, mở RQ2_SHAP rồi mở outputs. Mở README_results.md trước. Với CSV, có thể xem trực tiếp dưới dạng văn bản hoặc cài extension CSV/Excel viewer nếu muốn xem dạng bảng. Các file nên mở theo thứ tự: README_results.md, rq2_shap_top5_by_class.csv, rq2_shap_top5.csv, rq2_shap_sample_predictions.csv.")
    bullet(doc, "Muốn xem toàn lớp: mở rq2_shap_top5_by_class.csv và lọc true_label, is_correct = True, class_rank từ 1 đến 5.")
    bullet(doc, "Muốn xem một alert: mở rq2_shap_top5.csv và lọc sample_id. Xem rank, feature, shap_value và abs_shap_value.")
    bullet(doc, "Muốn tìm flow model dự đoán sai: mở rq2_shap_sample_predictions.csv và lọc is_correct = False, sau đó dùng sample_id đó tìm trong rq2_shap_top5.csv.")
    bullet(doc, "Không nên mở ranking_all_features.csv trước nếu máy chậm. File này có 106.314 dòng và chỉ cần khi phân tích đầy đủ 78 feature.")

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
