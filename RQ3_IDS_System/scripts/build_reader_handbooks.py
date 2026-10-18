"""Build two Vietnamese reader handbooks; no model/DB/network mutation.

Run with the bundled document Python runtime. Source snapshot: a680823,
2026-10-18. Generated Word files are editable content sources; their Word
layout is unverified. The separately typeset PDFs are visually checked.
"""
from pathlib import Path
import json
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'app'))
from feature_help import describe

OUT = ROOT / 'docs' / 'handbooks'
OUT.mkdir(exist_ok=True)


def base(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.78)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
        s = doc.styles[name]
        s.font.name = 'Arial'
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_after = Pt(7)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    doc.styles['Title'].font.size = Pt(25)
    doc.styles['Heading 1'].font.size = Pt(17)
    doc.styles['Heading 2'].font.size = Pt(13)
    doc.styles['Subtitle'].font.size = Pt(12)
    for name in ('Heading 1', 'Heading 2', 'Heading 3'):
        doc.styles[name].paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('RQ3  |  18/10/2026  |  ')
    r.font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.core_properties.author = 'Dự án RQ3 IDS'
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(subtitle, 'Subtitle')
    return doc


def p(d, text):
    d.add_paragraph(text)


def h(d, title):
    d.add_paragraph(title, 'Heading 2')


def chapter(d, title):
    d.add_page_break()
    d.add_paragraph(title, 'Heading 1')


def bullets(d, items):
    for text in items:
        d.add_paragraph(text, 'List Bullet')


def table(d, headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for c, width in zip(t.columns, widths): c.width = Inches(width)
    for c, value in zip(t.rows[0].cells, headers): c.text = value
    for row in rows:
        for c, value in zip(t.add_row().cells, row): c.text = str(value)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + side)
        for k, v in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]: el.set(qn('w:' + k), v)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + side); el.set(qn('w:w'), '90'); el.set(qn('w:type'), 'dxa'); margins.append(el)
    props.append(margins)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
        for j, cell in enumerate(row.cells):
            if widths: cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DCE6F1' if i == 0 else ('F5F7FA' if i % 2 == 0 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(sh)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(3)
                par.paragraph_format.line_spacing = 1.05
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.font.bold = i == 0
    d.add_paragraph().paragraph_format.space_after = Pt(0)


def code(d, text):
    for line in text.strip().splitlines():
        par = d.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.line_spacing = 1
        r = par.add_run(line)
        r.font.name = 'Consolas'; r.font.size = Pt(9)


def sources(d):
    chapter(d, 'Nguồn và cách kiểm chứng')
    p(d, 'Mốc nội dung: 18/10/2026, code nền a680823. Trạng thái là ảnh chụp tại thời điểm biên soạn, không phải cam kết hệ thống sẽ luôn có cùng số liệu. Đường dẫn nội bộ bên dưới tính từ thư mục SELF_KLTN. Các ví dụ ghi “minh họa” không phải kết quả đo.')
    for title, path in [
        ('P1 Kiến trúc và tiến độ', 'RQ3_IDS_System/docs/plan.md'),
        ('P2 Chạy hệ thống và đọc UI', 'RQ3_IDS_System/docs/dashboard_usage.md'),
        ('P3 Vận hành và bằng chứng tải', 'RQ3_IDS_System/docs/operations.md'),
        ('P4 Audit model và giải thích', 'RQ3_IDS_System/docs/xai_validation.md'),
        ('P5 Hợp đồng đặc trưng và audit cờ', 'RQ3_IDS_System/docs/feature_compatibility.md'),
        ('P6 Mã tính và diễn giải', 'RQ3_IDS_System/app/predict.py; normalize.py; explain_shap.py; explain_lime.py; reliability.py; feature_help.py'),
        ('P7 Bằng chứng PCAP và độ nhạy cờ', 'RQ3_IDS_System/docs/lifecycle_v3_final_evidence_20261018.json; flag_sensitivity_evidence_20261018.json'),
        ('P8 Bằng chứng Wi Fi', 'RQ3_IDS_System/docs/live_ht4_control_20261018.json; live_ht4_stop_retry_20261018.json; live_20min_ht3_retry_20261018.json'),
        ('P9 Huấn luyện và đánh giá gốc', 'RQ1_Model_Training_xboost/1_hyperparameter_tuning.py; 2_train_optimal_model.py'),
    ]:
        p(d, title + ': ' + path)
    p(d, 'Kiến thức nền tham khảo từ nguồn chính thức hoặc tác giả phương pháp, truy cập 18/10/2026. Các nguồn này giải thích khái niệm, không chứng nhận chất lượng dự án:')
    for s in [
        'NIST SP 800 94, Guide to Intrusion Detection and Prevention Systems: https://csrc.nist.gov/pubs/sp/800/94/final',
        'Microsoft, What is a security operations center: https://www.microsoft.com/en-us/security/business/security-101/what-is-a-security-operations-center-soc',
        'SHAP, phân biệt giải thích mô hình và quan hệ nhân quả: https://shap.readthedocs.io/en/latest/example_notebooks/overviews/Be%20careful%20when%20interpreting%20predictive%20models%20in%20search%20of%20causal%20insights.html',
        'Ribeiro, Singh và Guestrin, Why Should I Trust You, 2016: https://arxiv.org/abs/1602.04938',
    ]: p(d, s)


def guide():
    d = base('Hiểu và sử dụng hệ thống IDS RQ3', 'Quyển 1  Tổng quan SOC  Kiến trúc  Vận hành  Cách đọc kết quả')
    p(d, 'Hệ thống RQ3 là nguyên mẫu phát hiện xâm nhập dựa trên thống kê flow mạng và mô hình XGBoost. Hệ thống nhận CSV, trích xuất PCAP hoặc bắt traffic thật, dự đoán nhãn rồi hỗ trợ người phân tích bằng SHAP và LIME. Nó đã chạy được toàn đường, nhưng chưa được chứng nhận đủ chính xác và ổn định để bảo vệ mạng doanh nghiệp độc lập.')
    p(d, 'Tài liệu dành cho người trực tiếp chạy đề tài và người cần trình bày lại cho giảng viên. Đọc theo thứ tự nếu mới bắt đầu; khi thao tác, dùng chương 5 đến 8. Quyển 2 giải thích sâu các hạn chế và những kết luận không nên suy ra từ UI.')
    h(d, 'Bản đồ nội dung')
    bullets(d, ['1 IDS và vị trí trong SOC', '2 Hệ thống nhìn thấy và không nhìn thấy gì', '3 Kiến trúc và vòng đời một flow', '4 Những phần đã làm và phần còn thiếu', '5 Chuẩn bị và chạy bằng hai terminal', '6 Test realtime CSV và PCAP', '7 Đọc Dashboard và xử lý cảnh báo', '8 Đọc SHAP LIME và đối chứng', '9 Lỗi thường gặp và cách dừng an toàn', 'Phụ lục tra cứu đủ 78 đặc trưng và nguồn kiểm chứng'])
    h(d, 'Ba điều cần nhớ trước khi chạy')
    p(d, 'Một packet không phải một flow. Một cảnh báo model không phải sự cố đã xác nhận. Một lời giải thích đúng về cách model tính không bảo đảm quyết định ngoài thực tế là đúng.')
    p(d, 'Mốc biên soạn 18/10/2026. Không nhúng token, payload hoặc lịch sử IP cá nhân vào tài liệu. Không thay model hay cấu hình RQ2 khi biên soạn.')

    chapter(d, '1 IDS và vị trí trong SOC')
    p(d, 'IDS là Intrusion Detection System, hệ thống theo dõi dữ liệu và phát hiện dấu hiệu có thể liên quan xâm nhập. IPS còn có khả năng can thiệp để ngăn chặn. RQ3 hiện phát hiện và hiển thị cảnh báo; không tự chặn IP, ngắt kết nối hay cô lập máy. Phân biệt phát hiện với ngăn chặn giúp trình bày đúng phạm vi đề tài. [NIST; P1]')
    p(d, 'SOC là Security Operations Center, trung tâm hoặc chức năng vận hành an ninh gồm con người, quy trình và công cụ: nhận tín hiệu, phân loại, điều tra, ứng phó rồi rút kinh nghiệm. Không phải cứ có một dashboard là có một SOC. [Microsoft]')
    table(d, ['Thành phần', 'Vai trò', 'RQ3 hiện tại'], [
        ('NIDS hoặc cảm biến mạng', 'Quan sát traffic và phát hiện dấu hiệu theo mạng', 'Gần nhóm này: phân loại flow bằng ML'),
        ('HIDS hoặc EDR', 'Theo dõi hành vi trên host như tiến trình, file và hoạt động hệ điều hành', 'Không thu đầy đủ các loại telemetry này'),
        ('SIEM', 'Tập trung và tương quan nhiều nguồn log', 'Chưa có tích hợp SIEM được nghiệm thu'),
        ('SOAR và ứng phó', 'Tổ chức playbook, phối hợp hoặc tự động xử lý', 'Chưa có workflow ứng phó hoàn chỉnh'),
        ('Analyst', 'Đối chiếu chứng cứ và quyết định xử lý', 'UI và XAI hỗ trợ người dùng thực hiện việc này'),
    ], [1.45, 2.85, 2.6])
    h(d, 'Chạy trên laptop có phải endpoint security không')
    p(d, 'Laptop là nơi đặt cảm biến. Dữ liệu mà RQ3 phân tích vẫn là traffic mạng, không phải cây tiến trình, registry hay sự kiện file như EDR. Vì vậy nên mô tả là “cảm biến IDS theo flow chạy trên máy thử nghiệm”, không gọi là EDR chỉ vì chương trình chạy ở endpoint.')
    p(d, 'Vị trí phù hợp trong SOC là nguồn tín hiệu phát hiện và công cụ hỗ trợ phân tích ban đầu. Chuỗi đề xuất trong tương lai: cảm biến RQ3 → cảnh báo có metadata → SIEM hoặc analyst → điều tra → ứng phó được phê duyệt. Phần tích hợp đó là định hướng, chưa phải chức năng đã có.')

    chapter(d, '2 Hệ thống nhìn thấy và không nhìn thấy gì')
    p(d, 'Packet là một gói tin quan sát được. Flow gom các gói có liên quan theo địa chỉ, cổng, giao thức, chiều và chính sách thời gian. Một kết nối dài có thể bị chia thành nhiều flow; một flow không đồng nghĩa một người, một tab trình duyệt hoặc một lần tấn công.')
    h(d, 'Phạm vi bắt trên Wi Fi')
    p(d, 'Bắt interface Wi-Fi của laptop thường quan sát traffic đi/đến laptop và một số broadcast/multicast mà interface nhận được. Không được mặc định thấy toàn bộ traffic unicast của người khác trong cùng Wi-Fi. Muốn giám sát cả mạng phải thiết kế vị trí cảm biến, mirror/SPAN/TAP hoặc điểm trung chuyển phù hợp và có quyền quan sát. Đó chưa phải cấu hình đã nghiệm thu của đề tài.')
    p(d, 'Kết nối laptop vào hotspot điện thoại giúp tách môi trường test, nhưng Windows, ứng dụng nền, đồng bộ, cập nhật và các kết nối sẵn có vẫn tạo traffic. “Tôi chưa lướt web” không đồng nghĩa “không có packet”.')
    h(d, 'Có đọc được Facebook hay nội dung HTTPS không')
    p(d, 'Model nhận thống kê flow, không nhận nội dung hội thoại hay toàn bộ URL của trang web. Với traffic mã hóa, không thể suy ra nội dung ứng dụng chỉ từ cổng 443 hoặc IP. Cổng gợi ý dịch vụ chứ không xác nhận ứng dụng. Hệ thống hiện không có bước giải mã TLS để đọc nội dung.')
    h(d, 'Ba nguồn dữ liệu')
    table(d, ['Nguồn', 'Đầu vào và xử lý', 'Cần lưu ý'], [
        ('Realtime', 'Packet thật → extractor → CSV mới → feeder', 'Có độ trễ gom flow và cap tài nguyên'),
        ('PCAP', 'Packet đã lưu → trích xuất offline → gửi model', 'Không phát lại cuộc tấn công lên mạng khi chỉ trích file'),
        ('CSV', 'Đặc trưng đã có → chuẩn hóa → model', 'Có thể không có IP nếu file nguồn không chứa metadata'),
    ], [1, 3.0, 2.9])
    p(d, 'Metadata IP, thời gian, nguồn và mã lượt phục vụ truy vết. Không tự điền IP suy đoán khi thiếu. Không trộn “thời gian packet trong PCAP” với “thời gian hôm nay API xử lý file”. [P2, P5]')

    chapter(d, '3 Kiến trúc và vòng đời một flow')
    p(d, 'Đường dữ liệu giữ nguyên giữa cách chạy bốn terminal và cách dùng nút. Khác biệt là API điều khiển hai worker capture/feeder thay người dùng mở chúng bằng tay. Streamlit không giữ quyền sở hữu process. [P1, P3, P6]')
    table(d, ['Bước', 'Làm gì', 'Kết quả'], [
        ('1 Thu nhận', 'Bắt packet hoặc đọc PCAP; gom theo lifecycle v3', 'Một flow với thống kê và metadata'),
        ('2 Chuẩn hóa', 'Ánh xạ tên, kiểm tra đủ 78 cột, thứ tự, NaN/Inf và đơn vị', 'Vector model hợp lệ về cấu trúc'),
        ('3 Dự đoán', 'XGBoost trả 15 xác suất; chọn lớp lớn nhất', 'Nhãn, xác suất, model version'),
        ('4 Lưu và cảnh báo', 'Lưu flow/prediction; nhãn khác Benign tạo alert', 'Lịch sử truy vết trong SQLite'),
        ('5 Giải thích', 'SHAP nền cho alert; LIME khi có yêu cầu', 'Đóng góp và chẩn đoán lời giải thích'),
        ('6 Trình bày', 'Bộ lọc, biểu đồ, flow chi tiết, nhánh A/B', 'Analyst đọc và đối chiếu, chưa tự chặn'),
    ], [.85, 3.65, 2.4])
    p(d, '“Chuẩn hóa” ở đây chủ yếu là kiểm tra cấu trúc và quy đổi đã biết, không có nghĩa tự bảo đảm mọi công thức extractor giống train. IP nguồn/đích là metadata; Dst Port và Protocol có trong đầu vào model.')
    h(d, 'Trạng thái điều khiển')
    p(d, 'starting: đang khởi động; running: worker đã sẵn sàng; stopping: đang ngừng bắt/ghi nốt; draining: feeder gửi phần đã ghi; completed: chốt phiên theo đối soát; failed: lỗi hoặc phiên không đầy đủ; unknown: chưa xác minh chắc trạng thái process sau sự cố. Không coi unknown là đã dừng và không xóa journal để ép chạy mới.')
    p(d, 'Checkpoint và khóa idempotency giúp xử lý gửi lại sau mất ACK mà không tạo bản ghi trùng theo cùng khóa. Chúng không khôi phục packet chưa ghi ra CSV. Registry điều khiển phiên và DB dự đoán là hai loại lưu trữ khác nhau.')

    chapter(d, '4 Những phần đã làm và phần còn thiếu')
    table(d, ['Hạng mục', 'Đã có', 'Không được suy rộng thành'], [
        ('Dữ liệu và model', 'XGBoost RQ1, 78 cột, 15 lớp, model version', 'Model đúng trên mọi traffic live'),
        ('Nhập dữ liệu', 'CSV, PCAP, capture v3 và feeder checkpoint', 'Capture tương thích hoàn toàn extractor 2018'),
        ('XAI', 'TreeSHAP, LIME theo yêu cầu, guard suy biến', 'Giải thích là xác nhận tấn công'),
        ('UI', 'Lịch sử, chọn flow, biểu đồ, Realtime Start/Stop', 'SOC/EDR/SIEM hoàn chỉnh'),
        ('Vận hành', 'Cap, heartbeat, registry, stop/drain, backup/verify', 'Không mất dữ liệu trong mọi sự cố'),
        ('Kiểm thử', 'Fixture, API/model thật, AppTest, PCAP và smoke NIC', 'SLA tải cao hoặc chứng nhận production'),
    ], [1.15, 3, 2.75])
    h(d, 'Bằng chứng cần nhớ')
    p(d, 'Hai PCAP qua v3 có tổng 1.148 packet và 99 flow; 8 dự đoán Hulk. Không có ground truth đủ để lấy 8/99 làm FPR. SHAP của 8 cảnh báo đạt kiểm tra cộng dồn khoảng 0,00000286; LIME cả 8 thiếu neighborhood. [P4, P7]')
    p(d, 'Smoke điều khiển Wi-Fi 10,034 giây đã PASS với 267 packet và 19 flow CSV/API/ACK; drain thêm 2,133 giây. Một lượt dự kiến 30 giây trước đó FAIL khoảng giây 25 do cap. Lượt 20 phút riêng dùng cap thử 50.000 packet/64 MiB đã PASS ở tải quan sát; không thay mặc định 20.000/32 MiB và không chứng nhận mọi tải. [P3, P8]')
    p(d, 'Ngày 18/10 có chạy lại Stage14 trên DB tạm: model/API, CSV, PCAP, SHAP, LIME và UI AppTest PASS. Đây là nghiệm thu chức năng cô lập, không phải phép đo FPR mới hoặc kiểm thử trình duyệt/NIC dài hạn.')
    h(d, 'Còn phải hoàn thiện')
    p(d, 'Hợp đồng đầy đủ train–capture; đánh giá ground truth độc lập; benchmark độ trễ/tải dài hạn; kiểm thử browser thực; retention/log rotation; một số tình huống khôi phục process; kiểm chứng XAI toàn cohort và đóng gói triển khai. Đọc trạng thái mới nhất ở plan.md trước khi thuyết trình.')

    chapter(d, '5 Chuẩn bị và chạy bằng hai terminal')
    p(d, 'Mở VS Code tại RQ3_IDS_System. Dùng Python có dependencies API/UI, không dùng venv demo capture cũ. Capture có .venv_capture riêng và cần Npcap; quyền Administrator phụ thuộc cấu hình Npcap. Thiếu thư viện thì dùng requirements.txt trong đúng môi trường, không nâng phiên bản ngẫu nhiên. Cần có model, encoder, train_stats.json, lime_background.npy và dữ liệu tham chiếu theo config.yaml.')
    h(d, 'Terminal API')
    code(d, r'''cd C:\Users\LOQ\Downloads\SELF_KLTN\RQ3_IDS_System
Remove-Item Env:IDS_DB_PATH -ErrorAction SilentlyContinue
Remove-Item Env:IDS_CAPTURE_SESSION_DIR -ErrorAction SilentlyContinue
$env:IDS_CAPTURE_CONTROL_TOKEN = [guid]::NewGuid().ToString('N')
$env:IDS_CAPTURE_CONTROL_TOKEN
$env:IDS_CAPTURE_API_URL = 'http://127.0.0.1:8000'
python -m uvicorn app.main:app --host 127.0.0.1 --port 8000 --workers 1''')
    p(d, 'Hai lệnh Remove-Item chỉ bỏ biến override của terminal, không xóa file. Bôi đen token 32 ký tự và dùng Copy của menu chuột phải terminal. Ctrl+C có thể dừng chương trình; không dùng nó để copy. Đợi API khởi động hoàn tất. Không dùng --reload hoặc nhiều worker.')
    h(d, 'Terminal UI')
    code(d, r'''cd C:\Users\LOQ\Downloads\SELF_KLTN\RQ3_IDS_System
$env:IDS_CAPTURE_CONTROL_TOKEN = 'DAN_TOKEN_CUA_API'.Trim()
$env:IDS_API_URL = 'http://127.0.0.1:8000'
python -m streamlit run app/dashboard.py --server.address 127.0.0.1 --server.port 8501''')
    p(d, 'Thay DAN_TOKEN_CUA_API bằng token thực, giữ dấu nháy; .Trim() bỏ khoảng trắng thừa. Không tạo token khác ở terminal UI. Mở http://127.0.0.1:8501 trên chính máy. Không đưa UI control ra LAN/proxy/tunnel; token server-side không phải hệ thống đăng nhập cho người dùng từ xa.')
    p(d, 'Nếu đã có API/UI chạy, dừng phiên capture hợp tác trước rồi dừng dịch vụ cũ, không mở bản thứ hai cùng DB/registry. Các lệnh trên chỉ khởi động dịch vụ, chưa tự bắt Wi-Fi.')

    chapter(d, '6 Test realtime CSV và PCAP')
    h(d, 'Realtime thật')
    bullets(d, ['Vào Realtime, chọn đúng interface Wi-Fi; có thể kiểm tra tên bằng Get-NetAdapter.', 'Bấm Bắt đầu capture và đợi Đang bắt. Không khởi chạy thêm capture hoặc feeder bằng tay.', 'Lướt web trên máy. Quan sát packet nhận, flow xuất, ACK và heartbeat. Packet tăng nhưng flow chưa xuất có thể là đang gom flow.', 'Chọn Xem flow của phiên này rồi Giám sát flow. Kiểm tra sidebar không giữ bộ lọc nguồn/nhãn/thời gian loại mất dữ liệu.', 'Bấm Dừng capture. Đợi ghi nốt và gửi nốt hoàn thành rồi mới đóng API. Nếu failed, giữ log và xem lý do; không coi phiên đầy đủ.'])
    h(d, 'CSV đặc trưng')
    p(d, 'Nạp CSV / PCAP → CSV đặc trưng → chọn file → Gửi CSV vào /predict. File cần đủ 78 đặc trưng sau ánh xạ tên. Chỉ bật đổi giây sang micro giây nếu đúng loại CSV extractor; CSV dataset đã dùng micro giây không đổi lại. Kiểm tra số gửi thành công/lỗi và mã lượt. Thiếu IP thì UI ghi thiếu, không tự suy ra IP từ nhãn.')
    h(d, 'PCAP đã ghi')
    p(d, 'Nạp CSV / PCAP → PCAP đã ghi → chọn test.pcapng hoặc tesst2.pcapng → Trích xuất và gửi. Thao tác này đọc packet trong file, không tái phát packet lên mạng. Đợi trích xuất/gửi xong và chọn đúng lượt. Thời gian gói có thể là ngày cũ; API nhận là thời gian lần nhập mới.')
    h(d, 'Ghi nhận test có ích')
    p(d, 'Ghi thời điểm, thao tác, nguồn, mã lượt, extractor/model version, số flow và lỗi. Đừng chỉ chụp một nhãn đỏ. Nếu cần nghiên cứu báo nhầm, phải lưu ground truth có căn cứ và phạm vi chắc chắn; không gán mọi traffic nền là Benign chỉ vì người dùng không chủ động tấn công.')
    p(d, 'Sau reset DB, không resume checkpoint feeder cũ vì ACK cũ có thể làm bỏ qua các dòng đã bị xóa khỏi DB. Tạo lượt mới; UI capture tự tạo CSV/checkpoint riêng. Không cần reset DB giữa mọi lần thử. [P2, P3]')

    chapter(d, '7 Đọc Dashboard và xử lý cảnh báo')
    p(d, 'Đầu tiên chọn nguồn/lượt và khoảng thời gian. Các biểu đồ tổng quan theo bộ lọc không nhất thiết có cùng phạm vi với hàng đợi SHAP toàn hệ thống. Luôn đọc nhãn phạm vi trước khi so số liệu.')
    table(d, ['Thông tin', 'Đọc như thế nào'], [
        ('Flow mới nhất và phân trang', 'Mới nhất theo thời gian API xử lý; dùng nút trang và chọn dòng để mở chi tiết'),
        ('Nguồn và đích', 'IP:cổng nếu có metadata; không phải kết luận attacker/victim'),
        ('Nhãn và xác suất', 'Lớp lớn nhất trong 15 lớp; xác suất model, chưa chắc đã hiệu chuẩn'),
        ('SHAP', 'Chờ, đã có, bỏ/lỗi hoặc không cần cho Benign; không đồng nhất với trạng thái model'),
        ('Nhánh A hoặc B', 'Cách trình bày lời giải thích; B cần xem lại, không đổi nhãn'),
        ('Thời gian', 'Gói đầu, API nhận, dự đoán là các mốc khác nhau; UI hiển thị GMT+7'),
        ('Biểu đồ nhãn', 'Phân bố dự đoán trong phạm vi; không phải tỷ lệ tấn công thật'),
    ], [1.65, 5.25])
    h(d, 'Một quy trình đọc flow Hulk')
    p(d, 'Bước 1 xác nhận đúng lượt. Bước 2 đọc IP/cổng, giao thức và thời gian; tìm bối cảnh thao tác. Bước 3 xem xác suất mọi lớp để biết lớp đứng sau, không chỉ một số 96%. Bước 4 đọc giá trị đặc trưng và SHAP. Bước 5 xem cờ reliability và đối chứng. Bước 6 yêu cầu LIME nếu cần, đọc validity trước biểu đồ. Bước 7 ghi kết luận “cần kiểm tra” nếu chưa có chứng cứ độc lập.')
    p(d, 'Ví dụ: FIN đẩy điểm Hulk mạnh khi bạn duyệt web không chứng minh đã có DoS. Có thể là liên hệ model học hoặc khác biệt đặc trưng. Cần xem PCAP và định nghĩa extractor. Tương tự, nhãn Benign không chứng minh tuyệt đối an toàn và không được dùng làm nhãn thật để đánh giá chính model.')

    chapter(d, '8 Đọc SHAP LIME và đối chứng')
    h(d, 'SHAP là đóng góp cho điểm của model')
    p(d, 'Với lớp đang giải thích, SHAP dương đẩy điểm lên và SHAP âm kéo điểm xuống so với baseline của phép giải thích. Dự án dùng raw margin: phi không phải phần trăm xác suất. Kiểm tra cộng dồn dùng đủ 78 phi; Top-5 trên UI chỉ là phần dễ đọc. Các đặc trưng còn lại có thể đóng góp đáng kể. [P4, P6]')
    table(d, ['Ví dụ từ flow đã audit', 'Giá trị', 'SHAP gần đúng', 'Đọc đúng'], [
        ('FIN Flag Cnt', '1', '+5,865', 'Một dấu hiệu kết thúc TCP đẩy điểm Hulk lên'),
        ('Fwd Pkt Len Max', '1454', '+2,473', 'Độ dài gói góp tăng điểm Hulk trong flow này'),
        ('Dst Port', '443', '−0,483', 'Góp giảm điểm Hulk; không chứng minh an toàn'),
        ('Fwd IAT Min', '0', '−0,482', 'Góp giảm điểm Hulk; không phải luật chung'),
    ], [2.0, .65, 1.05, 3.2])
    h(d, 'LIME là mô hình đơn giản xấp xỉ quanh một flow')
    p(d, 'LIME tạo biến thể đầu vào, hỏi model và dùng trọng số theo khoảng cách để học lời giải thích cục bộ. Phải đọc trạng thái trước: insufficient_neighborhood nghĩa quá ít mẫu hữu dụng; zero_coefficients nghĩa hệ số đồng loạt bằng 0; undefined_fit nghĩa không có fit hữu hạn. ok chỉ vượt guard tối thiểu. Hệ số LIME không dùng cùng thang đo với SHAP raw margin; không so trực tiếp độ cao hai cột.')
    h(d, 'Nhánh và chỉ số hỗ trợ')
    p(d, 'Nhánh B bật nếu Top-5 có ít nhất 2 trong Dst Port, Fwd Seg Size Min, Init Fwd Win Byts hoặc lớp có dưới 100 mẫu train. Đây là cờ cần xem lại, không phải bằng chứng model sai. Nhánh A không phải chứng nhận model đúng. Proxy faithfulness thay Top-5 bằng median rồi hỏi lại model, chỉ là phép thử rút gọn; Jaccard@5 đo trùng tên đặc trưng, không đo độ đúng của nhãn.')
    p(d, 'Đối chứng so flow với 2.000 mẫu nền LIME và tập B nhỏ có nhãn. P05–P95 không phải ranh giới Benign/tấn công; thiếu mẫu một lớp thì phải đọc là chưa có dữ liệu. Độ lệch so nền không phải điểm nguy hiểm.')

    chapter(d, '9 Lỗi thường gặp và cách dừng an toàn')
    table(d, ['Hiện tượng', 'Cách xử lý'], [
        ('403 khi đọc trạng thái', 'Kiểm tra token API/UI giống nhau; bỏ khoảng trắng. Token API đọc lúc khởi động; không reset DB để sửa auth.'),
        ('Thông báo chỉ dùng local', 'Dùng bản đã sửa 99f5a03, bind 127.0.0.1 và mở tại máy. Không bỏ guard để công khai UI.'),
        ('Không có flow mới', 'Kiểm tra trạng thái capture, packet/flow/ACK, CSV, bộ lọc và lỗi. Đợi flow đóng/timeout; không suy ra từ API khỏe.'),
        ('Chạm cap', 'Phiên có lỗi/không đầy đủ. Giữ log, đối soát CSV/API; không tự bỏ cap hoặc đổi timeout để ép PASS.'),
        ('Unknown sau restart', 'Dùng Đối soát sau restart. Nếu chưa rõ, không kill PID hoặc xóa registry tùy ý.'),
        ('LIME không có biểu đồ', 'Đọc validity và ESS. Có thể là suy biến thật, không phải nút hỏng.'),
        ('Cảnh báo phiên bản model', 'Ghi phiên bản và đối chiếu môi trường; không tự nâng/hạ hoặc ghi đè artifact để hết cảnh báo.'),
    ], [1.8, 5.1])
    h(d, 'Kết thúc một lượt')
    p(d, 'Dừng bằng nút → đợi stopping/draining → kiểm tra completed hoặc failed → ghi mã lượt và kết quả → dừng UI rồi API bằng Ctrl+C khi không còn phiên active. Không đóng cửa sổ API giữa lúc ghi nốt nếu có thể tránh.')
    h(d, 'Backup và quyền truy cập')
    p(d, 'Dùng manage_storage.py theo operations.md để backup, verify và restore sang thư mục mới. Backup local không tự có trên GitHub. Không gửi PCAP chứa dữ liệu cá nhân lên Git nếu chưa rà soát. Chỉ bắt traffic trên máy/mạng được phép; token và bind local không thay thế phân quyền doanh nghiệp.')

    chapter(d, 'Phụ lục thuật ngữ dễ tra cứu')
    table(d, ['Thuật ngữ', 'Hiểu trong đề tài này'], [
        ('RQ1 RQ2 RQ3', 'RQ1 huấn luyện/đánh giá model; RQ2 nghiên cứu giải thích; RQ3 đưa model và XAI vào hệ thống chạy được. Runtime không tự train lại.'),
        ('API và UI', 'API nhận dữ liệu/tính/lưu; UI là giao diện người dùng xem và gửi lệnh.'),
        ('Extractor và feeder', 'Extractor tính thống kê từ packet; feeder chuyển từng dòng CSV sang API.'),
        ('Ground truth', 'Nhãn thật có căn cứ độc lập để kiểm tra dự đoán; không lấy từ chính model.'),
        ('Baseline và candidate', 'Bản mốc giữ nguyên để so sánh và bản thử nghiệm thay đổi có kiểm soát.'),
        ('Neighborhood và kernel', 'Các mẫu quanh flow theo thước đo; kernel quy khoảng cách thành trọng số mẫu gần/xa.'),
        ('ESS', 'Cỡ mẫu hiệu dụng theo trọng số; 1.000 mẫu sinh có thể chỉ tương đương khoảng một mẫu hữu dụng.'),
        ('OOD và drift', 'Đầu vào ngoài phân bố đã học và sự thay đổi phân bố theo môi trường/thời gian; chưa có bộ kiểm soát đầy đủ trong RQ3.'),
        ('ACK và checkpoint', 'Trong feeder: xác nhận API đã nhận và mốc tiến độ gửi. Khác cờ ACK của TCP trong đặc trưng mạng.'),
        ('RSS và backlog', 'Bộ nhớ tiến trình đang dùng và lượng công việc chưa xử lý xong.'),
        ('p50 p95 p99', 'Các phân vị: p95 là mốc mà 95% quan sát không vượt quá, với cohort và cách đo đã nêu.'),
        ('Fidelity và stability', 'Giải thích bám sát model đến đâu và thay đổi bao nhiêu khi lặp/thay đổi đầu vào; không đồng nghĩa nhãn đúng.'),
    ], [1.8, 5.1])
    stats = json.loads((ROOT/'data/train_stats.json').read_text(encoding='utf-8'))
    chapter(d, 'Phụ lục 15 nhãn và số mẫu huấn luyện')
    p(d, 'Số mẫu từ train_stats.json hiện có; không phải số flow đang chạy hay độ chính xác của lớp. Giữ nguyên cách viết nhãn của artifact, kể cả Infilteration. Nhãn khác Benign đều tạo cảnh báo theo logic hiện tại. Một nhãn có nhiều mẫu không bảo đảm tổng quát tốt sang traffic thật.')
    counts = stats['train_samples_per_class']
    assert len(counts) == 15 and sum(counts.values()) == 2800000
    table(d, ['Tên nhãn trong hệ thống', 'Số mẫu train'], [(name, f'{n:,}'.replace(',', '.')) for name,n in counts.items()], [4.9, 2])
    p(d, 'SQL Injection có 70 mẫu, dưới ngưỡng cờ lớp hiếm 100 đang cấu hình. Cờ chỉ nhắc analyst thận trọng; không làm đổi dự đoán hoặc tự khẳng định cảnh báo sai. DoS attacks-Hulk và các nhãn DDoS là các lớp riêng trong bảng, không gọi chung tất cả là Hulk.')
    names = list(stats['global_median'])
    assert len(names) == 78, len(names)
    for start in range(0, len(names), 10):
        chapter(d, f'Phụ lục đặc trưng {start + 1} đến {min(start + 10, 78)}')
        p(d, 'Ý nghĩa theo từ điển hiện tại của dự án, không phải chứng nhận tương đương dataset gốc. Fwd/Bwd là hai chiều của flow do extractor xác định; µs là micro giây ở đầu vào model. Giá trị riêng lẻ không có ngưỡng tấn công mặc định. [P5, P6]')
        rows = []
        for name in names[start:start+10]:
            info = describe(name)
            meaning = info['meaning']
            if info.get('note'): meaning += ' Lưu ý: ' + info['note']
            rows.append((name, info['unit'], meaning))
        table(d, ['Tên model', 'Đơn vị', 'Ý nghĩa và lưu ý'], rows, [1.55, 1.05, 4.3])
    sources(d)
    d.save(OUT/'01_Hieu_va_su_dung_he_thong_IDS_RQ3.docx')


def limitations():
    d = base('Hiểu các hạn chế của hệ thống IDS RQ3', 'Quyển 2  Model  Đặc trưng  SHAP  LIME  Độ trễ  Báo nhầm')
    p(d, 'Hệ thống đã hoạt động về chức năng nhưng còn khoảng cách giữa “chạy được” và “dự đoán đáng tin trên mạng thật”. Điểm yếu quan trọng nhất là chưa chứng nhận đầy đủ hợp đồng đặc trưng train–capture và chưa có đánh giá live độc lập có ground truth. SHAP giúp đọc quyết định của model; LIME hiện suy biến ở các cảnh báo PCAP đã audit. Cả hai không thay thế xác minh sự cố.')
    p(d, 'Quyển này trả lời các thắc mắc thường gặp bằng ví dụ, giải thích cơ chế và nêu cách kiểm chứng. Khi một nguyên nhân chưa được xác nhận, tài liệu ghi rõ là giả thuyết. Không dùng số cảnh báo trên Dashboard làm tỷ lệ báo nhầm.')
    h(d, 'Cách đọc mức độ chắc chắn')
    table(d, ['Mức', 'Nghĩa'], [
        ('Đã quan sát', 'Có code, audit hoặc kết quả kiểm thử xác định phạm vi'),
        ('Giả thuyết', 'Có khả năng giải thích hiện tượng nhưng chưa cô lập nguyên nhân'),
        ('Chưa đo', 'Chưa có phép đo thích hợp, không điền số ước đoán như kết quả'),
    ], [1.4, 5.5])
    h(d, 'Bản đồ nội dung')
    bullets(d, ['1 Model thật sự học và quyết định như thế nào', '2 Benign và Hulk tại sao có thể bị nhầm', '3 Khác biệt đặc trưng và chất lượng dữ liệu', '4 SHAP giải thích được gì và không chứng minh được gì', '5 LIME suy biến và giới hạn của cách lấy mẫu', '6 Cách đánh giá XAI và báo nhầm đúng', '7 Độ trễ từ packet đến giao diện', '8 Tải cao mất dữ liệu và giới hạn vận hành', '9 Hỏi đáp và kế hoạch hoàn thiện có điều kiện nghiệm thu'])

    chapter(d, '1 Model thật sự học và quyết định như thế nào')
    p(d, 'XGBoost là tập hợp nhiều cây quyết định. Trong huấn luyện, các cây học cách kết hợp các điều kiện trên con số đầu vào để giảm sai số dự đoán nhãn của dữ liệu đã có. Nó không được cung cấp một định nghĩa đầy đủ về “ý định tấn công”; nó học những liên hệ giúp phân biệt nhãn trong tập train.')
    p(d, 'Ví dụ minh họa: một cây hỏi một thống kê độ dài có nhỏ hơn một ngưỡng không, cây khác xét cờ hoặc khoảng cách gói; đóng góp của nhiều cây tạo điểm cho từng lớp. Đây không phải cây/ngưỡng thật được trích từ model. Muốn khẳng định model cụ thể học điều gì phải kiểm tra các split, ảnh hưởng trên cohort và SHAP, không tự viết một luật bằng lời thay cho model.')
    h(d, 'Quyết định của code hiện tại')
    p(d, 'Detector.predict_flow nhận đúng vector, lấy predict_proba và chọn chỉ số lớn nhất bằng argmax trong 15 lớp. main.py tạo alert nếu nhãn khác Benign. Hiện không có lớp “tôi không biết” hoặc bộ kiểm tra OOD đầy đủ để từ chối tất cả đầu vào lạ. Vì vậy model vẫn chọn một nhãn quen thuộc khi đầu vào khác xa train. [P6]')
    p(d, 'Một flow có xác suất Hulk 96% nghĩa model phân bổ xác suất như vậy trong hệ thống lớp của nó, không có nghĩa đã chứng minh 96% khả năng máy bị tấn công ngoài thực tế. Chưa có bằng chứng hiệu chuẩn xác suất trên môi trường live này.')
    h(d, 'Giới hạn của dữ liệu và đánh giá gốc')
    p(d, 'Model đọc train Parquet 2,8 triệu dòng và test 700 nghìn dòng, thuộc dữ liệu CIC-IDS2018 đã xử lý. Nhiều dòng không tự bảo đảm đại diện cho trình duyệt, giao thức, thiết bị và cấu hình mạng hiện tại. Mất cân bằng lớp hoặc mẫu lab đặc thù có thể làm ranh giới kém tổng quát; cần kiểm chứng theo lớp và theo môi trường.')
    p(d, 'Mã tuning RQ1 có đánh giá các cấu hình trên X_test. Vì tập này đã tham gia quá trình lựa chọn, không nên gọi nó là holdout hoàn toàn chưa từng đụng đến. Cần tập kiểm tra mới tách theo phiên/ngày/nguồn hợp lý; không lấy các flow live vừa dùng điều chỉnh để tự nghiệm thu. [P5, P9]')

    chapter(d, '2 Benign và Hulk tại sao có thể bị nhầm')
    p(d, 'Tên lớp trong hệ thống là DoS attacks-Hulk. Không đồng nhất với mọi dạng DDoS, và không khẳng định model nhận dạng đúng công cụ tạo traffic chỉ từ nhãn. Traffic bình thường bị gán Hulk là false positive chỉ khi có ground truth đủ tin cậy. Tấn công thật bị gán Benign là false negative; hiện chưa đủ dữ liệu có nhãn để lượng hóa cả hai trên môi trường của bạn.')
    h(d, 'Ba hướng nguyên nhân cần phân biệt')
    table(d, ['Hướng', 'Ví dụ dễ hiểu', 'Mức bằng chứng'], [
        ('Lệch cách biểu diễn', 'Tên cột giống nhưng tính SD, window, bulk hoặc chia flow khác', 'Có khác biệt audit; tác động từng flow cần cô lập'),
        ('Lệch môi trường', 'Traffic trình duyệt hiện tại khác các kịch bản lab train', 'Hợp lý, chưa đo đủ chất lượng live'),
        ('Model học dấu hiệu phụ', 'Header/cờ thường đi cùng một lớp trong dataset', 'SHAP gợi ý; không tự chứng minh quan hệ giả'),
    ], [1.45, 2.85, 2.6])
    h(d, 'Một giả thuyết đã bị bác bỏ trong phạm vi thử')
    p(d, 'Audit stage25 đổi cờ nguyên dương thành 1 trên bản sao vector của hai PCAP. 74/92 vector thay đổi nhưng 92/92 nhãn, xác suất và lá cây giữ nguyên. Model dùng split ngưỡng 1 ở các cờ đã kiểm tra, nên 1 và 18 cùng đi nhánh dương. Không có căn cứ nói chỉ ép ACK/FIN về nhị phân sẽ giảm Hulk ở những flow này. Kết luận không suy rộng sang model khác hoặc mọi phép sửa đầu vào. [P5, P7]')
    h(d, 'Bằng chứng chưa trả lời được điều gì')
    p(d, 'V3 tạo 99 flow từ hai PCAP, còn 8 nhãn Hulk; các sửa lifecycle không làm mất 8 cảnh báo đó. Chưa biết đầy đủ đóng góp nguyên nhân của từng sai khác đặc trưng vào nhãn. Cũng chưa có bộ nhãn ground truth để khẳng định tỷ lệ nhầm. Không nên train lại ngay chỉ để giảm màu đỏ trên UI.')
    p(d, 'Ví dụ minh họa: học phân biệt chó/mèo từ ảnh mà chó thường ở bãi cỏ có thể làm model dựa quá nhiều vào cỏ. Trong IDS, môi trường, kích thước header hoặc một cờ có thể đóng vai tương tự. SHAP chỉ ra model nhìn vào đâu; cần thí nghiệm mới kết luận đó có phải dấu hiệu phụ gây sai.')

    chapter(d, '3 Khác biệt đặc trưng và chất lượng dữ liệu')
    p(d, 'Một vector hợp lệ về cú pháp chưa chắc hợp lệ về ngữ nghĩa. Đủ 78 cột, không NaN/Inf và đúng thứ tự chỉ là bước đầu. Ví dụ cùng tên “độ dài gói” nhưng một bên tính frame, bên kia tính payload sẽ tạo số khác; đổi đơn vị thời gian đúng cũng không sửa được khác biệt công thức.')
    table(d, ['Nhóm', 'Quan sát trong audit', 'Hệ quả cần kiểm tra'], [
        ('Cờ TCP', 'Train 0/1; live có số đếm >1', 'Không đổi nhánh model trong audit cờ nhưng có thể ảnh hưởng LIME'),
        ('Bulk', 'Một số cột train toàn 0; flow có bulk lớn', 'Khoảng cách LIME rất lớn; chưa quy nhãn sai trực tiếp cho cột này'),
        ('Std độ dài', 'Train ở nhóm hai gói khớp sample SD, baseline dùng population', 'Cần thống nhất từng cột và miền, không đổi tất cả std theo suy đoán'),
        ('Ratio và window', 'Khác phép floor/chia thực hoặc sentinel -1/0', 'Có thể thay đường cây; giữ version và đối chiếu'),
        ('Mean và length', 'Quan hệ số học train khác một số công thức extractor', 'Cần nguồn lịch sử hoặc cặp PCAP–CSV chuẩn'),
        ('Lifecycle', 'V3 sửa timeout, FIN/ACK, SYN mới và packet trễ', 'Đúng chính sách mới chưa đồng nghĩa giống dataset 2018'),
    ], [1.15, 3, 2.75])
    p(d, 'Nguồn quan trọng còn thiếu là script tạo daily/train Parquet hoặc commit/config extractor lịch sử và cặp PCAP–CSV gốc. Java master hiện tại không tự đại diện phiên bản đã tạo dataset 2018. Các báo cáo Word nguồn không thay thế được công thức tái lập. [P5]')
    h(d, 'Cách xác minh đúng')
    p(d, 'Chọn nhóm đặc trưng, tạo packet nhỏ biết trước đáp án, tính tay hoặc bằng oracle độc lập, so baseline/candidate, rồi so vector/lá/xác suất trên PCAP thực. Lưu version, hash và sai số. Đừng cắt giá trị lớn, đổi timeout hoặc điền median chỉ để model bớt báo động. Dữ liệu thiếu không nên bị giả làm dữ liệu hợp lệ.')

    chapter(d, '4 SHAP giải thích được gì và không chứng minh được gì')
    p(d, 'TreeSHAP phân bổ chênh lệch giữa đầu ra model và baseline cho các đặc trưng của một flow. Trong RQ3, lớp được giải thích là lớp model vừa dự đoán, đầu ra là raw margin và UI giữ Top-5 theo trị tuyệt đối. Đây là lời giải thích cục bộ, không phải bảng xếp hạng tầm quan trọng toàn bộ model. [P6]')
    h(d, 'Vì sao FIN đứng đầu không có nghĩa FIN gây tấn công')
    p(d, 'FIN là dấu hiệu TCP kết thúc một chiều, xuất hiện hợp lệ trong kết nối bình thường. Khi FIN có SHAP dương cho Hulk, chỉ có thể nói nó góp phần đẩy điểm Hulk lên trong tổ hợp đầu vào này. Không được suy thành “FIN là bằng chứng DoS” hoặc “xóa FIN sẽ làm traffic an toàn”. Giải thích tương quan của model không phải quan hệ nhân quả. [SHAP]')
    h(d, 'Những giới hạn thật của SHAP')
    bullets(d, ['Giải thích model sai vẫn có thể chính xác về mặt tính toán; SHAP không tự phát hiện nhãn thật.', 'Kết quả phụ thuộc model và quy ước baseline/feature dependence. Các đặc trưng tương quan hoặc trùng thông tin làm cách phân bổ khó diễn giải như các nguyên nhân độc lập.', 'Top-5 không chứa toàn bộ 78 đóng góp. Không có tên trong Top-5 không nghĩa là tác động bằng 0.', 'Giá trị phi có dấu không phải phần trăm; không cộng Top-5 vào xác suất để tái tạo kết quả.', 'Một flow không đại diện mọi flow. Muốn kết luận model nhìn chung dựa vào gì cần cohort có phân lớp/nguồn và phân tích cả đúng lẫn sai.'])
    h(d, 'Bằng chứng hệ thống và giới hạn của bằng chứng')
    p(d, '8 cảnh báo PCAP audit có SHAP tái lập và sai số cộng dồn lớn nhất khoảng 0,00000286. Kiểm tra cộng dồn so baseline cộng đủ 78 phi với margin thật. Nó kiểm tra tính nhất quán, không kiểm tra tính đúng của nhãn, ngữ nghĩa extractor hay giá trị an ninh. [P4, P7]')
    p(d, 'Hai flow có FIN phi gần giống nhau có thể do cấu trúc cây và bối cảnh đường đi tương tự. Audit đã tái tạo các số này, không thấy gán cứng nhãn/phi. Từ điển định nghĩa và mẫu câu UI là nội dung tĩnh có chủ đích; giá trị, lớp, dấu phải đến từ kết quả flow.')

    chapter(d, '5 LIME suy biến và giới hạn của cách lấy mẫu')
    p(d, 'LIME hỏi “nếu thay đổi một số con số gần flow này, model phản ứng thế nào”, rồi học một mô hình đơn giản xấp xỉ phản ứng đó. Muốn xấp xỉ có nghĩa phải có đủ biến thể hữu ích và các mẫu phải đủ gần trong thước đo đang dùng. Chỉ sinh nhiều mẫu chưa bảo đảm điều này. [LIME; P4]')
    h(d, 'Tại sao 1000 mẫu mà gần như chỉ có 1 mẫu')
    p(d, 'Cấu hình khóa: background 2.000 mẫu, sinh 1.000 mẫu, kernel width 3,312, không discretize và mặc định lấy mẫu quanh nền thay vì quanh instance. Với flow lệch xa nền, kernel gán trọng số cực nhỏ cho mẫu sinh; mẫu gốc có khoảng cách 0 và trọng số 1. ESS gần 1 nghĩa hồi quy gần như chỉ nhìn một điểm, không đủ học ảnh hưởng thay đổi.')
    p(d, 'Flow đã audit có Bwd Blk Rate Avg khoảng 1,264 tỷ, trong khi toàn bộ nền ở cột đó bằng 0. Một số flow khác lệch ở ACK/FIN. Trong 8 cảnh báo, chỉ 2 flow có 999 weight bằng đúng 0; 6 flow còn lại có weight khác 0 nhưng rất nhỏ. Điểm chung là ESS xấp xỉ 1, không phải mọi flow đều underflow giống hệt nhau. [P4]')
    h(d, 'Tại sao R bình phương cao vẫn không dùng được')
    p(d, 'Một fit suy biến có thể trả R² bằng 1 hoặc 0 tùy trường hợp. Khi không có neighborhood hữu dụng, số đó không chứng minh chất lượng. Guard hiện tại đặt R²/Top-5/agreement không hợp lệ về null hoặc rỗng và giữ raw diagnostics, thay vì biến hệ số 0 thành một giải thích đẹp.')
    h(d, 'Các giới hạn khác')
    p(d, 'Cờ và cổng hiện được perturb như biến liên tục; các tổ hợp sinh có thể không hợp lệ với miền mạng. Sampling độc lập có thể phá quan hệ giữa tổng, trung bình và tốc độ. Kết quả phụ thuộc seed, nền, khoảng cách và kernel; fit tốt trên mẫu sinh chưa chắc trung thực trên traffic hợp lệ. Đổi seed nhiều lần rồi chọn biểu đồ đẹp sẽ tạo thiên lệch lựa chọn.')
    p(d, 'Không phải LIME bị cấm ở realtime: nó chạy theo yêu cầu trên flow đã hoàn tất. Cả 8 cảnh báo PCAP thử đều suy biến; 10/13 mẫu đối chứng nhỏ vượt guard tối thiểu. Không suy rộng hai nhóm nhỏ này thành tỷ lệ toàn hệ thống. Muốn cải thiện phải tách cấu hình thử nghiệm mới, giữ baseline RQ2 và đánh giá cả độ trung thực, ổn định, thời gian và hợp lệ miền.')

    chapter(d, '6 Cách đánh giá XAI và báo nhầm đúng')
    table(d, ['Chỉ số', 'Đo điều gì', 'Không chứng minh điều gì'], [
        ('SHAP additivity', 'Cộng dồn đóng góp khớp điểm model', 'Nhãn model đúng hoặc tấn công có thật'),
        ('LIME ESS', 'Số mẫu hữu dụng theo trọng số', 'ESS vừa >2 là giải thích tốt'),
        ('LIME R²', 'Fit mô hình xấp xỉ trên neighborhood', 'Độ chính xác IDS ngoài thực tế'),
        ('Jaccard Top 5', 'Mức trùng tập tên đặc trưng', 'Dấu/độ lớn giống nhau hoặc cả hai đúng'),
        ('Faithfulness proxy', 'Đổi Top-5 về median làm xác suất đổi bao nhiêu', 'Phép đo đầy đủ RQ2 hay quan hệ nhân quả'),
        ('Cờ môi trường', 'Top-5 có nhóm đặc trưng cần chú ý', 'Model chắc chắn học vẹt hoặc alert chắc chắn sai'),
    ], [1.4, 2.65, 2.85])
    h(d, 'Phân biệt tỷ lệ cảnh báo với FPR')
    p(d, 'FPR = FP / (FP + TN): trong các flow thực sự bình thường, bao nhiêu bị báo tấn công. FNR = FN / (FN + TP): trong các flow thực sự tấn công, bao nhiêu bị bỏ sót. Precision = TP / (TP + FP): trong cảnh báo model, bao nhiêu được xác nhận đúng. Nếu mẫu số bằng 0 thì chưa xác định, không ghi 0%.')
    p(d, 'Ví dụ minh họa: 1.000 flow bình thường đã xác minh, 40 bị báo tấn công → FPR 4%. Nếu chỉ biết “40 cảnh báo trên 1.000 flow chưa có nhãn” thì chỉ được nói tỷ lệ cảnh báo 4%. Flow không có ground truth phải tách riêng. Không dùng predicted_label làm label_true.')
    h(d, 'Không gán nhãn chỉ bằng đồng hồ')
    p(d, 'Trong khoảng có tấn công vẫn có traffic nền; một flow có thể cắt qua thời điểm bắt đầu/dừng. Ground truth nên dựa vào kịch bản được kiểm soát, tuple/dịch vụ và bằng chứng độc lập, với nhãn chưa rõ cho phần không chắc chắn. Bộ benign cần đa dạng thao tác nhưng biết phạm vi, không chỉ một lần mở Facebook.')
    p(d, 'Đánh giá mới phải có tập điều chỉnh riêng và tập kiểm tra chưa dùng chọn phương án. Báo confusion matrix, FPR/FNR theo lớp/nguồn, số mẫu và các flow loại trừ. Giảm số cảnh báo bằng cách tăng threshold có thể tăng bỏ sót; đó là trade-off cần đo chứ không phải tự động cải thiện.')

    chapter(d, '7 Độ trễ từ packet đến giao diện')
    p(d, 'Độ trễ phải nêu mốc đầu/cuối. Thời gian model suy luận không bao gồm thời gian chờ flow kết thúc; thời gian dừng/drain không phải thời gian phát hiện tấn công. PCAP phát lại có thời gian gói lịch sử, nên không lấy “API hôm nay trừ packet hôm trước” làm tốc độ xử lý.')
    table(d, ['Tầng', 'Nguồn chờ', 'Thông tin hiện có'], [
        ('Gom flow', 'Đóng TCP, timeout hoặc flush', 'Tuổi flow mặc định 90s; không phải SLA toàn đường'),
        ('CSV và feeder', 'Ghi file, poll, backlog, HTTP retry', 'Poll mặc định 1s; ACK mới cập nhật checkpoint'),
        ('API và model', 'Chuẩn hóa, suy luận, SQLite', 'Có timing thành phần; chưa báo p95 live mới'),
        ('SHAP nền', 'Đợi queue rồi tính và lưu', 'Cap queue 200 tác vụ đang chờ'),
        ('LIME theo yêu cầu', 'Nhận request, chờ lock và tính', 'Giới hạn nhận 4; instance tính tuần tự'),
        ('UI', 'Chu kỳ refresh, truy vấn và render', 'Panel capture refresh 3s; không phải cập nhật từng packet'),
    ], [1.1, 2.8, 3])
    h(d, 'Ví dụ chờ flow')
    p(d, 'Minh họa: packet đầu lúc 10:00:00; flow chỉ xuất lúc 10:01:30. Dù model tính rất nhanh, người dùng vẫn đã chờ khoảng 90 giây trước khi dữ liệu sẵn sàng dự đoán. Flow đóng sớm có thể xuất sớm hơn. Không nên hứa “phát hiện ngay khi tấn công bắt đầu” với kiến trúc flow hoàn tất này.')
    h(d, 'Cần đo tiếp')
    p(d, 'Ghi các mốc cùng clock hoặc có quy đổi rõ: packet đầu/cuối, flow export, feeder send/ACK, API nhận/predict, SHAP enqueue/start/end, UI hiển thị. Báo p50/p95/p99, max, throughput và backlog cho từng tải; tách flow khi còn sniff với flow chỉ xuất khi dừng. Không cộng các trung bình không cùng cohort rồi gọi là p95 toàn đường.')
    p(d, 'Lượt 10,034s đã đo drain 2,133s cho 19 flow ở tải quan sát. Đây là bằng chứng dừng hợp tác, không phải benchmark phát hiện. Chưa có số mới đủ để cam kết độ trễ tải cao. [P3, P8]')

    chapter(d, '8 Tải cao mất dữ liệu và giới hạn vận hành')
    p(d, 'Cap 20.000 là packet đang giữ đồng thời trong bộ đệm, không phải tổng packet cả lượt. Một phiên có thể xử lý nhiều hơn nếu flow được xuất và giải phóng kịp. Queue SHAP 200 là tác vụ giải thích, không phải packet. Hai tầng có cơ chế và hậu quả khác nhau.')
    table(d, ['Giới hạn hoặc sự cố', 'Hậu quả thực tế'], [
        ('Buffer capture đầy', 'Dừng lỗi/phiên không đầy đủ; không coi dữ liệu đã hiện là toàn bộ traffic'),
        ('SHAP queue đầy', 'Dự đoán/cảnh báo vẫn lưu; giải thích có thể bị bỏ lượt và chờ recovery'),
        ('Capture chết trước xuất CSV', 'Packet trong RAM chưa bền; checkpoint feeder không khôi phục phần này'),
        ('Mất ACK sau commit', 'Idempotency/checkpoint hỗ trợ retry không trùng theo khóa; không bảo đảm NIC không rơi gói'),
        ('Registry unknown', 'Chặn khởi động trùng khi chưa rõ process; cần đối soát thay vì xóa khóa'),
        ('Disk hoặc log tăng dài hạn', 'Có cảnh báo/guard nhưng chưa có retention và rotation hoàn chỉnh'),
    ], [2.0, 4.9])
    p(d, 'Giới hạn mặc định còn có 4.096 flow mở và 32 MiB byte packet được tính theo cơ chế của capture. Byte cap không phải RSS tổng tiến trình Python. Bản lưu packet gọn giảm payload/object nhưng không bỏ giới hạn số packet. Nâng cap chỉ tăng chỗ đệm, không sửa được một hệ thống xử lý chậm hơn đầu vào mãi mãi.')
    h(d, 'Bằng chứng tải phải đọc đúng phạm vi')
    p(d, 'Lượt HT4 30s FAIL khoảng 25,1s vì cap, với 20.056 packet nhận/xuất, 57 flow xuất nhưng chỉ 7 ACK trước abort. Lượt ngắn sau đó PASS không xóa bằng chứng thất bại. Lượt 20 phút khác với cap thử 50.000/64 MiB PASS nhưng peak buffered 48.495, gần giới hạn; không phải lý do bỏ cap mặc định. [P3, P8]')
    p(d, 'Đối soát accepted/exported chỉ trong phạm vi extractor đã nhận. Mất gói ở NIC/Npcap, đường mạng hoặc vị trí không nhìn thấy không được suy ra bằng 0. Cảm biến hiện không là TCP stack/reassembly đầy đủ; chưa bao phủ mọi giao thức như một sản phẩm NIDS hoàn chỉnh.')

    chapter(d, '9 Hỏi đáp và kế hoạch hoàn thiện')
    h(d, 'Train thêm dữ liệu có giải quyết được không')
    p(d, 'Có thể giúp nếu nguyên nhân là thiếu đại diện, nhưng phải kiểm tra đầu vào trước. Thêm dữ liệu với công thức lệch hoặc nhãn suy đoán có thể làm sai thêm và phá so sánh RQ1/RQ2. Giữ model baseline bất biến; model mới là phiên bản/thí nghiệm riêng với holdout độc lập.')
    h(d, 'Đổi timeout nhỏ hơn để thấy kết quả nhanh hơn được không')
    p(d, 'Có thể làm flow xuất sớm hơn nhưng thay đổi số gói, duration, IAT và các thống kê. Đây là thay đổi đầu vào model, không chỉ tối ưu tốc độ UI. Phải version, đo tác động lên nhãn/XAI và độ trễ, không chỉnh âm thầm.')
    h(d, 'LIME hỏng thì SHAP có dùng được không')
    p(d, 'Có thể tiếp tục dùng SHAP như thông tin về model nếu kiểm tra tính toán đạt. Không suy từ LIME hỏng sang SHAP sai; cũng không suy từ SHAP cộng dồn đúng sang model đúng. UI cần giữ hai trạng thái độc lập.')
    h(d, 'Có thể tự chặn IP dựa trên Hulk chưa')
    p(d, 'Chưa nên. Chưa có ground truth đủ, chưa hiệu chuẩn theo môi trường và có nguy cơ báo nhầm; chặn tự động có thể ngắt dịch vụ hợp lệ. Phạm vi hiện tại là hỗ trợ quan sát và analyst, không tự ứng phó.')
    h(d, 'Thứ tự hoàn thiện đề xuất')
    bullets(d, ['1 Chốt môi trường tái lập model/encoder và hợp đồng đặc trưng bằng nguồn/fixture. Nghiệm thu bằng vector và xác suất, không chỉ số cảnh báo giảm.', '2 Xây bộ benign/attack có nhãn xác minh, chia điều chỉnh/kiểm tra độc lập và đo FPR/FNR theo nguồn, lớp và phiên.', '3 Đo latency từng tầng, backlog/RSS và tối ưu capture giữ đúng ngữ nghĩa; soak test tải khác nhau và sự cố dừng/restart.', '4 Thử LIME cải tiến trong cấu hình mới; đánh giá hợp lệ miền, fidelity, stability, ESS và thời gian so với baseline khóa.', '5 QA trình duyệt và thao tác người dùng, giảm lỗi token/đọc trạng thái; bổ sung retention, triển khai và tích hợp SOC chỉ khi có yêu cầu.'])
    p(d, 'Mỗi bước phải ghi mục đích, điều kiện nghiệm thu, kết quả và giới hạn vào tài liệu tương ứng rồi commit riêng. Không đánh dấu hoàn thành chỉ vì test fixture PASS hoặc giao diện đẹp hơn.')
    sources(d)
    d.save(OUT/'02_Han_che_model_SHAP_LIME_va_realtime_RQ3.docx')


if __name__ == '__main__':
    guide()
    limitations()
    print('Built two reader handbooks in', OUT)
