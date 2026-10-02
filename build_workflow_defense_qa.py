from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'RQ2_Tai_lieu' / 'Bo_cau_hoi_va_kich_ban_bao_ve_Workflow_RQ1_RQ2_RQ3.docx'

def font(run, size=10.5, bold=False, color=(25,25,25)):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor(*color)

def p(doc, text='', size=10.5, bold=False, after=5):
    x=doc.add_paragraph(); x.paragraph_format.space_after=Pt(after); x.paragraph_format.line_spacing=1.13
    font(x.add_run(text), size, bold); return x

def h(doc, text, level=1):
    x=doc.add_paragraph(style=f'Heading {level}'); x.paragraph_format.space_before=Pt(14 if level==1 else 9); x.paragraph_format.space_after=Pt(5); x.paragraph_format.keep_with_next=True
    font(x.add_run(text), 15 if level==1 else 12, True, (0,0,0)); return x

def qa(doc, q, a):
    x=doc.add_paragraph(); x.paragraph_format.space_before=Pt(4); x.paragraph_format.space_after=Pt(2); x.paragraph_format.line_spacing=1.12
    font(x.add_run('Hỏi: '+q), 10.5, True, (30,65,110))
    y=doc.add_paragraph(); y.paragraph_format.left_indent=Inches(.18); y.paragraph_format.space_after=Pt(4); y.paragraph_format.line_spacing=1.12
    font(y.add_run('Đáp: '+a))

def quote(doc, text):
    x=doc.add_paragraph(); x.paragraph_format.left_indent=Inches(.2); x.paragraph_format.right_indent=Inches(.15); x.paragraph_format.space_before=Pt(5); x.paragraph_format.space_after=Pt(8)
    font(x.add_run('Trả lời ngắn khi bảo vệ: '), 10.5, True, (30,65,110)); font(x.add_run('“'+text+'”'))

def block(doc, n, name, purpose, qs):
    h(doc, f'{n}. {name}', 2); p(doc, purpose)
    for q,a in qs: qa(doc,q,a)

def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
    title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_after=Pt(8); font(title.add_run('Bộ câu hỏi và kịch bản bảo vệ workflow RQ1 RQ2 RQ3'),20,True,(0,0,0))
    sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after=Pt(15); font(sub.add_run('Ôn tập mô hình IDS, SHAP LIME và demo realtime'),10.5,False,(80,80,80))
    p(doc,'Tài liệu tổng hợp các câu hỏi và đáp án đã trao đổi: workflow từ CICIDS2018 đến dashboard, lý do chọn Random Forest/XGBoost, SHAP/LIME, ba tiêu chí RQ2 và giới hạn realtime của mô hình flow based.', after=10)

    h(doc,'A. Kịch bản thuyết trình workflow',1)
    for t in [
        'Workflow gồm ba phần. RQ1 chọn model phát hiện tốt nhất; RQ2 kiểm chứng lời giải thích của model; RQ3 đưa model và lời giải thích vào dashboard giám sát traffic thật.',
        'Đầu vào là CSE-CIC-IDS2018. Mỗi dòng là một network flow có nhãn Benign hoặc tấn công. Dữ liệu được làm sạch, chuẩn hóa nhãn và giữ đúng 78 feature mạng trước khi tách train và test.',
        'Trong RQ1, em train Random Forest và XGBoost, sau đó so sánh trên cùng test set bằng Accuracy, Macro Precision, Macro Recall, Macro F1, kết quả từng lớp và confusion matrix. XGBoost là model tối ưu được cố định cho RQ2 và RQ3.',
        'Trong RQ2, em khóa cohort test đại diện. SHAP và LIME cùng giải thích một flow, cùng model XGBoost và cùng lớp dự đoán. Em đánh giá Agreement, Domain Validation và Faithfulness để kiểm tra lời giải thích.',
        'Trong RQ3, packet thật được gom thành flow, trích thành 78 feature tương thích với train set, đưa vào XGBoost để dự đoán. Explainer được chọn từ RQ2 trả về Top feature và dashboard hiển thị cảnh báo kèm lý do.'
    ]: p(doc,t,after=7)
    quote(doc,'RQ1 trả lời model nào phát hiện tốt; RQ2 trả lời lời giải thích có đáng tin không; RQ3 chứng minh model và explanation có thể dùng trong giám sát flow mạng gần thời gian thực.')

    doc.add_page_break(); h(doc,'B. Câu hỏi theo từng khối workflow',1)
    blocks=[
    ('1','CSE CIC IDS2018 Dataset','Dataset đầu vào gồm các flow mạng đã có nhãn.',[
      ('Vì sao chọn CSE-CIC-IDS2018?','Đây là dataset IDS phổ biến, có Benign và nhiều loại tấn công; feature flow phù hợp để train IDS.'),('Một dòng dữ liệu là gì?','Là một network flow, tức phiên trao đổi giữa nguồn và đích, không phải một packet đơn.'),('Đề tài dùng bao nhiêu lớp?','15 lớp gồm Benign và 14 lớp tấn công.')]),
    ('2','Data Preprocessing','Làm sạch dữ liệu và tạo schema đầu vào ổn định.',[
      ('Vì sao phải tiền xử lý?','Để xử lý NaN, Infinity, dữ liệu không hợp lệ và nhãn không thống nhất.'),('78 feature là gì?','Là các đặc trưng flow như duration, packet length, IAT, TCP flags và port.'),('Vì sao RQ3 phải giữ đúng 78 feature?','Vì XGBoost đã học đúng schema đó; sai tên hoặc thứ tự có thể làm dự đoán sai.')]),
    ('3','Train Test Split','Tách dữ liệu để đánh giá khách quan.',[
      ('Vì sao chia train và test?','Train để model học, test để đánh giá trên dữ liệu chưa thấy.'),('Đánh giá trên train set có được không?','Không đáng tin vì kết quả cao có thể do model nhớ dữ liệu.'),('Vì sao giữ nguyên test set?','Để so sánh các model công bằng và tránh rò rỉ dữ liệu.')]),
    ('4','RQ1 Train Random Forest and XGBoost','Hai mô hình cây được train trên cùng 78 feature.',[
      ('Vì sao chọn Random Forest và XGBoost?','Cả hai hợp dữ liệu bảng phi tuyến; Random Forest là baseline ổn định, XGBoost là boosting mạnh.'),('Vì sao không chọn một model ngay từ đầu?','Hai model giúp so sánh có căn cứ thay vì chọn cảm tính.'),('Model nhận và trả về gì?','Nhận 78 feature của flow và trả về một trong 15 nhãn.')]),
    ('5','Evaluate and Compare Models','Đánh giá trên cùng test set bằng metric phù hợp dữ liệu mất cân bằng.',[
      ('Dùng chỉ số nào?','Accuracy, Macro Precision, Macro Recall, Macro F1, per-class metrics, confusion matrix và FPR Benign.'),('Vì sao không chỉ Accuracy?','Benign thường nhiều; Accuracy cao có thể che việc bỏ sót lớp tấn công hiếm.'),('Metric quyết định chính?','Macro F1 vì cân bằng Precision và Recall cho các lớp.')]),
    ('6','Select Optimal Model','Khóa model tốt nhất cho RQ2 và RQ3.',[
      ('Model tối ưu là gì?','XGBoost với 100 cây, depth 8, learning rate 0,08 và tree method hist.'),('Vì sao không chọn Random Forest?','XGBoost có Macro F1 và hiệu năng tổng thể cao hơn trên test set.'),('Dùng model này ở đâu?','Để giải thích ở RQ2 và dự đoán realtime ở RQ3.')]),
    ('7','RQ2 Select Test Samples','Khóa cohort để SHAP và LIME giải thích cùng flow.',[
      ('Vì sao không chạy LIME trên 700 nghìn flow?','LIME phải tạo nhiều mẫu lân cận và gọi model nhiều lần cho từng flow nên rất tốn tài nguyên.'),('Cohort ghép cặp là gì?','Tập flow test cố định, đại diện, có cả kết quả SHAP và LIME.'),('Vì sao cùng cohort?','Để so sánh công bằng: cùng flow, model và lớp dự đoán.')]),
    ('8','Run SHAP','SHAP phân bổ đóng góp feature cho dự đoán XGBoost.',[
      ('SHAP giải thích gì?','Mức và chiều đóng góp của từng feature cho dự đoán.'),('SHAP dương và âm nghĩa gì?','Dương đẩy về lớp xét; âm kéo dự đoán rời lớp đó.'),('SHAP có local/global không?','Có; local cho từng flow, global tổng hợp nhiều flow.')]),
    ('9','Run LIME','LIME xấp xỉ quyết định quanh một flow bằng mô hình tuyến tính nhỏ.',[
      ('LIME hoạt động thế nào?','Nó tạo mẫu lân cận, quan sát output XGBoost rồi fit mô hình tuyến tính local.'),('LIME là global hay local?','Chủ yếu là local.'),('LIME khác SHAP thế nào?','LIME là xấp xỉ local; SHAP dựa trên Shapley values nên hai ranking có thể khác.')]),
    ('10','Compare SHAP and LIME','Kiểm tra độ tin cậy của lời giải thích.',[
      ('Mục tiêu so sánh?','Xem hai phương pháp có tương tự, hợp lý và phản ánh model hay không.'),('Có so trực tiếp độ lớn SHAP với LIME không?','Không, vì thang đo khác nhau; so Top-k, thứ hạng và chiều ảnh hưởng.'),('Khác nhau có nghĩa là sai?','Không; cần kiểm tra tiếp domain validation và faithfulness.')]),
    ('11','Agreement Do the Top Features Match','Đo mức SHAP và LIME cùng chọn feature quan trọng.',[
      ('Agreement là gì?','Mức hai phương pháp chọn feature giống nhau trên cùng flow.'),('Jaccard at k là gì?','Tỷ lệ feature trùng nhau giữa Top-k SHAP và LIME.'),('Spearman dùng làm gì?','Đo độ tương đồng về thứ hạng feature.')]),
    ('12','Domain Validation Do Features Match Attack Mechanisms','Đối chiếu ranking với kiến thức an ninh mạng.',[
      ('Domain validation là gì?','Kiểm tra feature quan trọng có hợp lý với cơ chế tấn công không.'),('Ví dụ SSH Bruteforce?','Dst Port, số packet, packet rate và IAT có thể là dấu hiệu hợp lý.'),('Nếu feature không hợp lý thì sao?','Model có thể học dấu vết dataset thay vì cơ chế tấn công tổng quát.')]),
    ('13','Faithfulness Do Features Truly Affect the Model Decision','Kiểm tra trực tiếp feature có tác động đến model không.',[
      ('Faithfulness là gì?','Mức lời giải thích phản ánh hành vi thật của model.'),('Kiểm tra bằng cách nào?','Thay Top-k feature bằng giá trị tham chiếu rồi so mức giảm confidence với thay ngẫu nhiên.'),('Nếu confidence giảm mạnh hơn random?','Đó là bằng chứng feature được giải thích thực sự ảnh hưởng XGBoost.')]),
    ('14','RQ2 Visualizations and Report','Chuyển kết quả thô thành bằng chứng nghiên cứu.',[
      ('Đầu vào là gì?','CSV SHAP, LIME, agreement, domain, stability, robustness và faithfulness.'),('Đầu ra là gì?','Bảng, biểu đồ và báo cáo RQ2.'),('Vì sao không chỉ CSV?','CSV khó nhìn xu hướng; biểu đồ giúp trình bày và bảo vệ.')]),
    ('15','Select Explainer for Dashboard','Chọn explainer chính cho luồng realtime.',[
      ('Vì sao phải chọn một explainer chính?','Chạy cả hai cho mọi flow làm chậm và có thể gây mâu thuẫn cho người dùng.'),('Có thể vẫn dùng cả hai không?','Có; một phương pháp realtime mặc định, phương pháp kia để đối chiếu.'),('Chọn dựa vào đâu?','Faithfulness, stability, domain validation, tốc độ và độ phù hợp với XGBoost.')]),
    ('16','RQ3 Real Time Network Traffic','Nhận traffic thật ở dạng packet.',[
      ('Khác CICIDS2018 thế nào?','CICIDS2018 là flow đã có nhãn; realtime là packet thật chưa có nhãn.'),('Traffic đến từ đâu?','Website demo, máy người dùng hoặc môi trường lab.'),('RQ3 chứng minh gì?','Model và explainer có thể tích hợp vào giám sát thực tế, không chỉ chạy CSV offline.')]),
    ('17','Feature Extraction','Biến packet thành vector 78 feature.',[
      ('Vì sao không đưa packet thô vào model?','Model train theo feature cấp flow, không phải packet.'),('CICFlowMeter làm gì?','Gom packet thành flow và tính duration, packet length, IAT, flags, port, tốc độ.'),('Rủi ro lớn nhất?','Feature realtime sai tên, thiếu cột hoặc sai thứ tự schema train.')]),
    ('18','Prediction Optimal Model','XGBoost dự đoán từ vector feature chuẩn hóa.',[
      ('Dùng model nào?','File pkl của XGBoost optimal model.'),('Đầu vào gì?','Đúng 78 feature theo đúng thứ tự train.'),('Đầu ra gì?','Nhãn dự đoán và confidence.')]),
    ('19','Explanation Selected Explainer','Giải thích dự đoán của flow cụ thể.',[
      ('Explainer nhận gì?','Feature flow vừa dự đoán và model XGBoost.'),('Trả về gì?','Top feature, giá trị và chiều ảnh hưởng.'),('Vì sao chạy sau prediction?','Phải biết model dự đoán lớp nào rồi mới giải thích lý do cho lớp đó.')]),
    ('20','IDS Dashboard Prediction Alert and Explanation','Hiển thị kết quả cho người dùng.',[
      ('Dashboard hiển thị gì?','Thời gian, nguồn, đích, port, protocol, nhãn, confidence và Top feature.'),('Dashboard có phải model không?','Không; nó là giao diện, XGBoost là thành phần dự đoán.'),('Giá trị thực tế?','Vừa cảnh báo vừa cho biết lý do, giảm tính black box.')]),
    ]
    for args in blocks: block(doc,*args)

    doc.add_page_break(); h(doc,'C. Câu hỏi về flow packet và realtime',1)
    more=[
      ('Một lần truy cập web sinh rất nhiều packet, xử lý sao?','Không dự đoán từng packet. Packet liên quan được gom thành flow rồi mới tính feature và dự đoán.'),
      ('Packet nào thuộc cùng flow?','Thường gom theo 5 tuple: source IP, destination IP, source port, destination port và protocol; theo dõi cả hai chiều.'),
      ('Một lần bấm web chỉ tạo một flow?','Không. Có thể có DNS, web server, API, CDN. Dashboard hiển thị flow; có thể gom các flow đáng ngờ thành alert.'),
      ('Khi nào flow được xuất?','Khi FIN/RST, idle timeout hoặc active timeout tùy pipeline.'),
      ('Có phát hiện ngay packet đầu tiên không?','Không với model flow based. Hệ thống cần đủ packet để tính feature.'),
      ('Vậy vì sao gọi realtime?','Realtime nghĩa là xử lý liên tục với độ trễ hữu hạn đủ để phản ứng; với đề tài này là near real time ở mức flow/window.'),
      ('Giảm trễ như thế nào?','Xuất snapshot theo active timeout hoặc time window, ví dụ vài giây, rồi dự đoán khi flow còn tiếp diễn.'),
      ('Snapshot có hạn chế gì?','Model train trên full flow, còn snapshot là partial flow nên nên đánh giá thêm hoặc train theo window nếu triển khai nghiêm túc.'),
      ('Độ trễ gồm gì?','Chờ flow/window, trích feature, XGBoost prediction, explanation và dashboard hiển thị.'),
      ('SQL Injection/XSS có phát hiện trực tiếp payload không?','Không trực tiếp với 78 flow feature; cần WAF hoặc mô hình phân tích HTTP payload riêng.')]
    for q,a in more: qa(doc,q,a)
    quote(doc,'Hệ thống là flow based near real time IDS: packet được bắt liên tục, gom thành flow hoặc snapshot, trích 78 feature tương thích CICIDS2018 rồi đưa vào XGBoost. Cảnh báo không ở packet đầu tiên nhưng có độ trễ đo và kiểm soát được.')
    h(doc,'D. Câu hỏi đề xuất để hỏi giảng viên hướng dẫn',1)
    p(doc,'Thầy ơi, vì mô hình của em được train trên CICIDS2018 với 78 feature do CICFlowMeter tạo ra, sang RQ3 em nên triển khai realtime theo hướng nào để feature từ traffic thật tương thích với feature lúc train? Em có cần dùng đúng CICFlowMeter, cùng phiên bản/cấu hình hay cần mapping và xử lý trung gian trước khi đưa vào model pkl?')
    h(doc,'E. Checklist trước khi demo RQ3',1)
    for x in ['Khóa file pkl, danh sách và thứ tự 78 feature.','Xử lý NaN, Infinity, cột thừa/thiếu trước prediction.','Đo latency từ flow/window đến dashboard.','Hiển thị prediction trước, explanation sau nếu cần.','Chỉ demo trên môi trường lab và traffic hợp pháp.']:
        y=doc.add_paragraph(style='List Bullet'); font(y.add_run(x))
    doc.core_properties.title='Bo cau hoi va kich ban bao ve workflow RQ1 RQ2 RQ3'
    doc.save(OUT); print(OUT)
if __name__=='__main__': main()
